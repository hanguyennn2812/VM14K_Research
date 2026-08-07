#!/usr/bin/env python3
"""
build_topic_review_xlsx.py — turn topic_mapping_proposed.csv into a workbook a
doctor can fill in.

Read-only with respect to the dataset. Reads
reports/analysis/topic_mapping_proposed.csv and writes
reports/analysis/topic_mapping_review.xlsx with three sheets:

    HuongDan   instructions + live progress counters
    Review     one row per raw topic string, with reviewer dropdowns
    DanhMuc    the paper's 34 canonical specialties (dropdown source)

    python scripts/analysis/build_topic_review_xlsx.py
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from propose_topic_mapping import CANONICAL_34

CSV_PATH = REPO_ROOT / "reports" / "analysis" / "topic_mapping_proposed.csv"
OUT_PATH = REPO_ROOT / "reports" / "analysis" / "topic_mapping_review.xlsx"

FONT = "Arial"

# Reviewer choices
DECISIONS = [
    "Đồng ý",
    "Sửa - xem cột H",
    "Loại bỏ (không dùng topic này)",
    "Cần thảo luận",
]
EXTRA_TARGETS = [
    "KHONG THUOC 34 - de nghi them chuyen khoa moi",
    "LOAI BO",
]

TAG_FILL = {
    "AMBIGUOUS": "FFF2CC",  # amber - needs judgement
    "JUNK": "F2F2F2",       # grey  - not a specialty
    "AUTO": "DDEBF7",       # blue  - close variant, low risk
    "MATCH": "E2EFDA",      # green - already canonical
}
TAG_VN = {
    "AMBIGUOUS": "CAN QUYET DINH",
    "JUNK": "KHONG PHAI CHUYEN KHOA",
    "AUTO": "BIEN THE GAN",
    "MATCH": "DA CHUAN",
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_rows() -> list[dict]:
    with CSV_PATH.open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def style_header(ws, row: int, ncols: int, fill: str = "1F4E79") -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_danhmuc(wb: Workbook):
    ws = wb.create_sheet("DanhMuc")
    ws["A1"] = "34 chuyên khoa chuẩn (VM14K paper, Table 5 / Appendix A)"
    ws["A1"].font = Font(name=FONT, bold=True, size=11)
    ws["A2"] = "Danh sách này là nguồn cho dropdown ở cột H sheet Review. Không sửa."
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="808080")

    ws["A4"] = "Chuyên khoa"
    style_header(ws, 4, 1)
    for i, name in enumerate(CANONICAL_34 + EXTRA_TARGETS, start=5):
        cell = ws.cell(row=i, column=1, value=name)
        cell.font = Font(name=FONT, size=10)
        cell.border = BORDER
        if name in EXTRA_TARGETS:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.column_dimensions["A"].width = 46
    ws.sheet_view.showGridLines = False
    return ws, 5, 5 + len(CANONICAL_34) + len(EXTRA_TARGETS) - 1


def build_review(wb: Workbook, rows: list[dict], dm_first: int, dm_last: int):
    ws = wb.create_sheet("Review")
    headers = [
        "STT",
        "Chuỗi topic gốc\n(trong dữ liệu)",
        "Số dòng\nảnh hưởng",
        "Phân loại\nđề xuất",
        "Chuyên khoa\nđề xuất",
        "Lý do (nhóm phân tích)",
        "► QUYẾT ĐỊNH\n(chọn)",
        "► Chuyên khoa đúng\n(nếu sửa)",
        "► Ghi chú của bác sĩ",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 42

    for i, r in enumerate(rows, start=2):
        tag = r["decision_tag"]
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=r["raw_string"])
        ws.cell(row=i, column=3, value=int(r["frequency"]))
        ws.cell(row=i, column=4, value=TAG_VN[tag])
        ws.cell(row=i, column=5, value=r["proposed_canonical"])
        ws.cell(row=i, column=6, value=r["note"])
        for c in range(1, 10):
            cell = ws.cell(row=i, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(c in (2, 5, 6, 9)),
                horizontal="center" if c in (1, 3, 4) else "left",
            )
            if c <= 6:
                cell.fill = PatternFill("solid", fgColor=TAG_FILL[tag])
            else:
                # reviewer input cells: yellow, per workbook-for-filling-in convention
                cell.fill = PatternFill("solid", fgColor="FFFFCC")
        ws.cell(row=i, column=3).number_format = "#,##0"

    last = len(rows) + 1

    dv_decision = DataValidation(
        type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True
    )
    dv_decision.error = "Chọn một giá trị trong danh sách."
    dv_decision.errorTitle = "Giá trị không hợp lệ"
    ws.add_data_validation(dv_decision)
    dv_decision.add(f"G2:G{last}")

    dv_target = DataValidation(
        type="list",
        formula1=f"DanhMuc!$A${dm_first}:$A${dm_last}",
        allow_blank=True,
    )
    dv_target.error = "Chọn một chuyên khoa trong danh mục 34 (sheet DanhMuc)."
    dv_target.errorTitle = "Giá trị không hợp lệ"
    ws.add_data_validation(dv_target)
    dv_target.add(f"H2:H{last}")

    # highlight any row still undecided
    ws.conditional_formatting.add(
        f"G2:G{last}",
        FormulaRule(formula=[f'$G2=""'], fill=PatternFill("solid", fgColor="FCE4EC")),
    )

    widths = {"A": 6, "B": 30, "C": 10, "D": 20, "E": 26, "F": 52, "G": 22, "H": 30, "I": 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{last}"
    ws.sheet_view.showGridLines = False
    return ws, last


def build_huongdan(wb: Workbook, rows: list[dict], last_row: int):
    ws = wb.create_sheet("HuongDan", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for col in "CDE":
        ws.column_dimensions[col].width = 20

    def put(cell, text, *, bold=False, size=10, color="000000", italic=False):
        ws[cell] = text
        ws[cell].font = Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
        ws[cell].alignment = Alignment(vertical="top", wrap_text=True)

    put("B2", "Chuẩn hóa danh mục chuyên khoa VM14K — phiếu duyệt", bold=True, size=14)
    put(
        "B4",
        "Dữ liệu VM14K hiện có 127 chuỗi topic khác nhau, trong khi paper công bố là 34 "
        "chuyên khoa. Nhóm phân tích đã đề xuất sẵn cách xử lý cho từng chuỗi. "
        "Việc của bác sĩ: xác nhận hoặc sửa lại đề xuất đó.",
        size=10,
    )
    put("B6", "CÁCH LÀM — chỉ điền 3 cột màu vàng ở sheet 'Review'", bold=True, size=11)
    put("B7", "Cột G — QUYẾT ĐỊNH", bold=True)
    put("C7", "Chọn từ dropdown: Đồng ý / Sửa / Loại bỏ / Cần thảo luận")
    put("B8", "Cột H — Chuyên khoa đúng", bold=True)
    put("C8", "CHỈ điền khi cột G = 'Sửa'. Chọn từ dropdown (34 chuyên khoa).")
    put("B9", "Cột I — Ghi chú", bold=True)
    put("C9", "Tùy chọn. Ghi lý do nếu quyết định khác đề xuất.")
    put(
        "B11",
        "Không cần sửa các cột A–F (nền màu) — đó là dữ liệu và đề xuất của nhóm phân tích.",
        italic=True,
        color="808080",
    )

    put("B13", "VÍ DỤ (dòng mẫu, không có trong sheet Review)", bold=True, size=11)
    ex_head = ["Chuỗi gốc", "Đề xuất", "► Cột G", "► Cột H", "► Cột I"]
    for j, h in enumerate(ex_head):
        ws.cell(row=14, column=2 + j, value=h)
    style_header(ws, 14, 6)
    ex = ["Hepatology", "Gastroenterology", "Sửa - xem cột H", "Internal Medicine",
          "Gan mật nên xếp Nội khoa"]
    for j, v in enumerate(ex):
        c = ws.cell(row=15, column=2 + j, value=v)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor="FFFFCC" if j >= 2 else "F2F2F2")
    ws.row_dimensions[15].height = 30

    put("B17", "Ý NGHĨA CÁC NHÓM PHÂN LOẠI", bold=True, size=11)
    legend = [
        ("DA CHUAN", "Đã đúng 1 trong 34 chuyên khoa. Không cần làm gì.", "MATCH"),
        ("BIEN THE GAN", "Chỉ khác cách viết / là chuyên khoa con rõ ràng. Rủi ro thấp.", "AUTO"),
        ("CAN QUYET DINH", "Là lĩnh vực y học thật nhưng KHÔNG nằm trong 34. Cần bác sĩ quyết.", "AMBIGUOUS"),
        ("KHONG PHAI CHUYEN KHOA", "Từ chỉ quy trình / cơ quan / quản lý. Đề nghị bỏ.", "JUNK"),
    ]
    r = 18
    for name, desc, tag in legend:
        c = ws.cell(row=r, column=2, value=name)
        c.font = Font(name=FONT, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=TAG_FILL[tag])
        c.border = BORDER
        c.alignment = Alignment(vertical="center")
        d = ws.cell(row=r, column=3, value=desc)
        d.font = Font(name=FONT, size=10)
        d.alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        r += 1

    put("B24", "TIẾN ĐỘ (tự cập nhật khi điền)", bold=True, size=11)
    prog = [
        ("Tổng số chuỗi cần xem", f"=COUNTA(Review!$B$2:$B${last_row})"),
        ("Đã có quyết định", f'=COUNTIF(Review!$G$2:$G${last_row},"<>")'),
        ("Còn trống", "=C26-C27"),
        ("Đồng ý", f'=COUNTIF(Review!$G$2:$G${last_row},"{DECISIONS[0]}")'),
        ("Sửa", f'=COUNTIF(Review!$G$2:$G${last_row},"{DECISIONS[1]}")'),
        ("Loại bỏ", f'=COUNTIF(Review!$G$2:$G${last_row},"{DECISIONS[2]}")'),
        ("Cần thảo luận", f'=COUNTIF(Review!$G$2:$G${last_row},"{DECISIONS[3]}")'),
    ]
    r = 26
    for label, formula in prog:
        lc = ws.cell(row=r, column=2, value=label)
        lc.font = Font(name=FONT, size=10)
        lc.alignment = Alignment(vertical="center")
        vc = ws.cell(row=r, column=3, value=formula)
        vc.font = Font(name=FONT, bold=True, size=10)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border = BORDER
        vc.fill = PatternFill("solid", fgColor="E2EFDA")
        r += 1

    put(
        "B35",
        "Nguồn: reports/analysis/topic_mapping_proposed.csv — sinh bởi "
        "scripts/analysis/propose_topic_mapping.py. Danh mục 34 chuyên khoa lấy nguyên văn "
        "từ paper VM14K (arXiv 2506.01305) Table 5 / Appendix A. "
        "Dữ liệu gốc KHÔNG bị thay đổi bởi bước này.",
        size=9,
        color="808080",
        italic=True,
    )
    return ws


def main() -> None:
    rows = load_rows()
    wb = Workbook()
    wb.remove(wb.active)

    _, dm_first, dm_last = build_danhmuc(wb)
    _, last_row = build_review(wb, rows, dm_first, dm_last)
    build_huongdan(wb, rows, last_row)

    wb._sheets = [wb["HuongDan"], wb["Review"], wb["DanhMuc"]]
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH.relative_to(REPO_ROOT)}  ({len(rows)} rows to review)")


if __name__ == "__main__":
    main()
