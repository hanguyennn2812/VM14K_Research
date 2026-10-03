#!/usr/bin/env python3
"""
build_review_workbook.py — collect the test questions where models disagree
with the answer key into one workbook a reviewer can fill in.

Read-only with respect to the dataset and the runs. Reads the zero-shot runs in
reports/eval/runs/ and, when present, the explanation logs written by
collect_reasoning.py in reports/eval/reasoning/. Writes to reports/eval/review/:

    VM14K_review.xlsx    the workbook (sheets below)
    review_ids.txt       ids of BatDong + ToanBoSai, input for collect_reasoning.py
    reasoning_log.md     every question with each model's explanation, for reading

Sheets:
    HuongDan       how each list is built, columns, caveats
    BatDong        DeepSeek-v4.1-flash and Gemma-4-31B pick the same non-key option
    ToanBoSai      all 7 fully-run models miss the key
    DoKho          accuracy per difficulty label, per model
    TatCa          every test question with every model's letter (for filtering)
    LapLuan        one row per (question, model) explanation, full text

    python scripts/eval/build_review_workbook.py
"""
from __future__ import annotations

import collections
import glob
import json
import os
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_PATH = REPO_ROOT / "data" / "cleaned" / "clean_final.jsonl"
RUNS_DIR = REPO_ROOT / "reports" / "eval" / "runs"
REASONING_DIR = REPO_ROOT / "reports" / "eval" / "reasoning"
OUT_DIR = REPO_ROOT / "reports" / "eval" / "review"

LETTERS = "ABCDEFG"
FONT = "Arial"
EXCEL_CELL_MAX = 32000  # Excel's hard limit is 32,767 characters

# (column label, run file stem, counts toward the all-models-wrong set).
# Nemotron-3-Ultra stopped at 1,431/1,658 questions, so it is shown but not counted.
MODELS = [
    ("DeepSeek-v4.1-flash", "nvidia__deepseek-ai_deepseek-v4.1-flash__paper__test", True),
    ("Gemma-4-31B", "nvidia__google_gemma-4-31b-it__paper__test", True),
    ("Nemotron-3-Ultra", "nvidia__nvidia_nemotron-3-ultra-550b-a55b__paper__test", False),
    ("Gemma-4-12B", "gemma4_12b__paper__test", True),
    ("Qwen3.5-9B", "qwen3.5_9b__paper__test", True),
    ("Qwen3-8B", "qwen3_8b__paper__test", True),
    ("Llama-3.1-8B", "llama3.1_8b__paper__test", True),
    ("MedGemma-4B", "medgemma_4b__paper__test", True),
]
STRONG_PAIR = ("DeepSeek-v4.1-flash", "Gemma-4-31B")
THIRD = "Nemotron-3-Ultra"
COUNTED = [label for label, _, counted in MODELS if counted]

# Friendly names for the explanation logs, keyed by the API model id.
REASONING_LABELS = {
    "deepseek-ai/deepseek-v4.1-flash": "DeepSeek-v4.1-flash",
    "google/gemma-4-31b-it": "Gemma-4-31B",
    "nvidia/nemotron-3-ultra-550b-a55b": "Nemotron-3-Ultra",
    "gemma4:12b": "Gemma-4-12B",
    "qwen3.5:9b": "Qwen3.5-9B",
    "qwen3:8b": "Qwen3-8B",
    "llama3.1:8b": "Llama-3.1-8B",
    "medgemma:4b": "MedGemma-4B",
    "openai/gpt-oss-20b": "gpt-oss-20b",
}

REVIEW_CATEGORIES = [
    "Đáp án chuẩn sai",
    "Câu hỏi lỗi / mơ hồ / thiếu dữ kiện",
    "Đúng theo phác đồ / thực hành Việt Nam",
    "Kiến thức cũ / hướng dẫn đã thay đổi",
    "Model sai, đáp án chuẩn đúng",
    "Không chắc / cần thảo luận",
]

FILL_OK = PatternFill("solid", fgColor="E2EFDA")
FILL_BAD = PatternFill("solid", fgColor="F8CBAD")
FILL_NONE = PatternFill("solid", fgColor="EDEDED")
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_REVIEW = PatternFill("solid", fgColor="FFF2CC")
FONT_HEAD = Font(name=FONT, bold=True, color="FFFFFF")
FONT_BODY = Font(name=FONT, size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def load_rows():
    with open(DATA_PATH, encoding="utf-8") as fh:
        return {r["id"]: r for r in (json.loads(line) for line in fh if line.strip())}


def load_run(stem):
    """{id: last error-free record} — the same rule run_eval.py uses to resume."""
    out = {}
    with open(RUNS_DIR / f"{stem}.jsonl", encoding="utf-8") as fh:
        for line in fh:
            if not line.strip():
                continue
            rec = json.loads(line)
            if rec.get("error") is None:
                if rec.get("perm") != list(range(rec["n_options"])):
                    raise ValueError(f"{stem}: {rec['id']} was asked with shuffled options")
                out[rec["id"]] = rec
    return out


def load_reasoning():
    """({label: {id: record}}, {label: model label in MODELS}) from collect_reasoning.py logs.

    One entry per log file, last error-free row per id. A file is one run configuration
    (provider, model, think), so files are never merged: when a model has several logs,
    each label names its provider and thinking setting."""
    found = []
    for path in sorted(glob.glob(str(REASONING_DIR / "*.jsonl"))):
        recs = {}
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                if line.strip():
                    rec = json.loads(line)
                    if rec.get("error") is None:
                        recs[rec["id"]] = rec
        if recs:
            meta = next(iter(recs.values()))
            found.append((Path(path).stem, REASONING_LABELS.get(meta["model"], meta["model"]), meta, recs))
    per_base = collections.Counter(base for _, base, _, _ in found)
    logs, base_of = {}, {}
    for stem, base, meta, recs in found:
        label = base
        if per_base[base] > 1:
            label = f"{base} [{meta['provider']}, think {'on' if meta['think'] else 'off'}]"
        if label in logs:
            label = f"{label} {stem}"
        logs[label], base_of[label] = recs, base
    return logs, base_of


def load_annotations(path):
    """{sheet: {id: reviewer values}} already typed into an existing workbook, so a rebuild keeps them."""
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True)
    out = {}
    for name in ("BatDong", "ToanBoSai"):
        if name not in wb.sheetnames:
            continue
        rows_iter = wb[name].iter_rows(values_only=True)
        header = list(next(rows_iter, []))
        missing = [h for h in ["id"] + REVIEW_HEADERS if h not in header]
        if missing:
            raise SystemExit(f"{path.name}/{name}: columns {missing} not found — renamed by hand? "
                             "Refusing to overwrite reviewer input.")
        id_col = header.index("id")
        cols = [header.index(h) for h in REVIEW_HEADERS]
        for r in rows_iter:
            values = [r[c] if c < len(r) else None for c in cols]
            if r and r[id_col] and any(v not in (None, "") for v in values):
                out.setdefault(name, {})[r[id_col]] = values
    wb.close()
    return out


def options_block(row):
    return "\n".join(f"{LETTERS[i]}. {opt}" for i, opt in enumerate(row["options"]))


def option_text(row, letter):
    if not letter:
        return ""
    i = LETTERS.index(letter)
    return row["options"][i] if i < len(row["options"]) else ""


def clip(text):
    text = text or ""
    if len(text) > EXCEL_CELL_MAX:
        return text[:EXCEL_CELL_MAX] + "\n…(cắt bớt, bản đầy đủ trong reports/eval/reasoning/)"
    return text


def style_header(ws, headers, widths):
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = FONT_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"


def fill_letter(cell, letter, gold):
    cell.alignment = Alignment(horizontal="center", vertical="top")
    cell.fill = FILL_NONE if not letter else (FILL_OK if letter == gold else FILL_BAD)


def add_review_columns(ws, first_col, n_rows):
    """Reviewer columns with dropdowns, shaded so they stand out."""
    cat = DataValidation(type="list", formula1='"' + ",".join(REVIEW_CATEGORIES) + '"', allow_blank=True)
    letter = DataValidation(type="list", formula1='"' + ",".join(LETTERS) + '"', allow_blank=True)
    ws.add_data_validation(cat)
    ws.add_data_validation(letter)
    last = n_rows + 1
    cat.add(f"{get_column_letter(first_col)}2:{get_column_letter(first_col)}{last}")
    letter.add(f"{get_column_letter(first_col + 1)}2:{get_column_letter(first_col + 1)}{last}")
    for r in range(2, last + 1):
        for c in range(first_col, first_col + 5):
            cell = ws.cell(row=r, column=c)
            cell.fill = FILL_REVIEW
            cell.alignment = WRAP_TOP


REVIEW_HEADERS = ["Phân loại (người duyệt)", "Đáp án đúng theo người duyệt",
                  "Nguồn / bằng chứng", "Ghi chú", "Người duyệt"]
REVIEW_WIDTHS = [30, 14, 40, 40, 14]


def write_question_sheet(ws, ids, rows, runs, reasoning, base_of, extra_cols, annotations):
    """Shared layout for BatDong and ToanBoSai. extra_cols: [(header, width, fn(id)->value)];
    annotations: {id: reviewer values} carried over from the previous workbook."""
    reasoning_labels = list(reasoning)
    headers = (["STT", "id", "Chuyên khoa", "Độ khó", "Câu hỏi", "Các phương án",
                "Đáp án chuẩn", "Nội dung đáp án chuẩn"]
               + [h for h, _, _ in extra_cols]
               + [f"{label}" for label, _, _ in MODELS]
               + [f"Số model đúng (/{len(COUNTED)})"]
               + [h for label in reasoning_labels
                  for h in (f"{label} — đáp án khi giải thích", f"{label} — giải thích")]
               + REVIEW_HEADERS)
    widths = ([5, 12, 18, 11, 50, 55, 8, 30]
              + [w for _, w, _ in extra_cols]
              + [9] * len(MODELS) + [9]
              + [w for _ in reasoning_labels for w in (12, 70)]
              + REVIEW_WIDTHS)
    style_header(ws, headers, widths)
    model_col0 = 9 + len(extra_cols)
    for n, qid in enumerate(ids, start=1):
        row = rows[qid]
        gold = row["answer"]
        letters = [runs[label].get(qid, {}).get("pred") for label, _, _ in MODELS]
        n_ok = sum(runs[label].get(qid, {}).get("correct", False) for label in COUNTED)
        values = ([n, qid, ", ".join(row["medical_topic"]), row["difficulty_level"], row["question"],
                   options_block(row), gold, option_text(row, gold)]
                  + [fn(qid) for _, _, fn in extra_cols]
                  + [l or "" for l in letters] + [n_ok])
        for label in reasoning_labels:
            rec = reasoning[label].get(qid)
            if rec is None:
                values += ["", ""]
                continue
            orig = runs.get(base_of[label], {}).get(qid, {}).get("pred")
            flag = "" if not orig or rec["pred"] == orig else f" (đổi từ {orig})"
            values += [f"{rec['pred'] or 'Không chọn'}{flag}", clip(rec.get("answer_text"))]
        values += annotations.get(qid, [""] * len(REVIEW_HEADERS))
        ws.append(values)
        r = ws.max_row
        for cell in ws[r]:
            cell.font = FONT_BODY
            cell.alignment = WRAP_TOP
        for i, letter in enumerate(letters):
            fill_letter(ws.cell(row=r, column=model_col0 + i), letter, gold)
        col = model_col0 + len(MODELS) + 1
        for label in reasoning_labels:
            rec = reasoning[label].get(qid)
            if rec is not None:
                fill_letter(ws.cell(row=r, column=col), rec["pred"], gold)
                ws.cell(row=r, column=col + 1).alignment = TOP  # long text: no wrap, read in formula bar
            col += 2
    add_review_columns(ws, len(headers) - len(REVIEW_HEADERS) + 1, len(ids))
    ws.auto_filter.ref = ws.dimensions


def main():
    rows = load_rows()
    runs = {label: load_run(stem) for label, stem, _ in MODELS}
    reasoning, base_of = load_reasoning()
    workbook_path = OUT_DIR / "VM14K_review.xlsx"
    annotations = load_annotations(workbook_path)
    test_ids = sorted(runs[STRONG_PAIR[0]])

    a, b = (runs[m] for m in STRONG_PAIR)
    batdong = [i for i in test_ids
               if i in b and a[i]["pred"] and a[i]["pred"] == b[i]["pred"] != a[i]["gold"]]
    toanbosai = [i for i in test_ids
                 if all(i in runs[m] for m in COUNTED) and not any(runs[m][i]["correct"] for m in COUNTED)]
    batdong_set, toanbosai_set = set(batdong), set(toanbosai)

    def third_status(qid):
        rec = runs[THIRD].get(qid)
        if rec is None:
            return "Chưa trả lời"
        return "Có" if rec["pred"] == a[qid]["pred"] else f"Không (chọn {rec['pred'] or '?'})"

    def majority(qid):
        """(every letter with the top vote count, that count) — ties are kept, not broken."""
        votes = collections.Counter(runs[m][qid]["pred"] for m in COUNTED if runs[m][qid]["pred"])
        if not votes:
            return [], 0
        top = max(votes.values())
        return sorted(letter for letter, n in votes.items() if n == top), top

    batdong.sort(key=lambda i: (third_status(i) != "Có", third_status(i) == "Chưa trả lời",
                                sum(runs[m][i]["correct"] for m in COUNTED)))
    toanbosai.sort(key=lambda i: -majority(i)[1])

    wb = Workbook()
    guide = wb.active
    guide.title = "HuongDan"

    ws = wb.create_sheet("BatDong")
    write_question_sheet(ws, batdong, rows, runs, reasoning, base_of, [
        ("Đáp án 2 model mạnh cùng chọn", 10, lambda i: a[i]["pred"]),
        ("Nội dung đáp án đó", 30, lambda i: option_text(rows[i], a[i]["pred"])),
        ("Nemotron-3-Ultra cũng chọn?", 14, third_status),
        ("Cũng nằm trong ToanBoSai?", 10, lambda i: "Có" if i in toanbosai_set else ""),
    ], annotations.get("BatDong", {}))

    ws = wb.create_sheet("ToanBoSai")
    write_question_sheet(ws, toanbosai, rows, runs, reasoning, base_of, [
        ("Đáp án được chọn nhiều nhất", 10,
         lambda i: " / ".join(majority(i)[0]) + (" (hòa)" if len(majority(i)[0]) > 1 else "")),
        (f"Số model chọn (/{len(COUNTED)})", 9, lambda i: majority(i)[1]),
        ("Nội dung đáp án đó", 30,
         lambda i: "\n".join(f"{l}. {option_text(rows[i], l)}" for l in majority(i)[0])),
        ("Cũng nằm trong BatDong?", 10, lambda i: "Có" if i in batdong_set else ""),
    ], annotations.get("ToanBoSai", {}))

    # Reviewer input whose question left both lists (definitions changed) is kept, not dropped.
    listed = {"BatDong": batdong_set, "ToanBoSai": toanbosai_set}
    orphans = [(sheet, qid, values) for sheet, by_id in annotations.items()
               for qid, values in by_id.items() if qid not in listed[sheet]]
    if orphans:
        ws = wb.create_sheet("ChuThichCu")
        style_header(ws, ["Sheet cũ", "id"] + REVIEW_HEADERS, [12, 34] + REVIEW_WIDTHS)
        for sheet, qid, values in orphans:
            ws.append([sheet, qid] + values)

    # Difficulty: accuracy per label for every model, plus the mean over the counted ones.
    ws = wb.create_sheet("DoKho")
    levels = ["Easy", "Medium", "Challenging", "Hard"]
    style_header(ws, ["Độ khó (nhãn)", "Số câu"] + [label for label, _, _ in MODELS]
                 + [f"Trung bình {len(COUNTED)} model"], [16, 9] + [12] * len(MODELS) + [14])
    ws.freeze_panes = "B2"
    by_level = collections.defaultdict(list)
    for qid in test_ids:
        by_level[rows[qid]["difficulty_level"]].append(qid)
    for level in levels:
        ids = by_level[level]
        accs = []
        for label, _, _ in MODELS:
            answered = [i for i in ids if i in runs[label]]
            accs.append(100 * sum(runs[label][i]["correct"] for i in answered) / len(answered)
                        if answered else None)
        mean = sum(accs[i] for i, (_, _, c) in enumerate(MODELS) if c) / len(COUNTED)
        ws.append([level, len(ids)] + [None if a is None else round(a, 1) for a in accs] + [round(mean, 1)])
        for cell in ws[ws.max_row]:
            cell.font = FONT_BODY
    ws.append([])
    ws.append(["Độ chính xác (%) trên tập test. Nemotron-3-Ultra chỉ chạy 1.431/1.658 câu nên không tính vào trung bình. "
               "Nhãn Hard chỉ có 13 câu, sai số rất lớn."])

    ws = wb.create_sheet("TatCa")
    headers = (["id", "Chuyên khoa", "Độ khó", "Câu hỏi", "Đáp án chuẩn"]
               + [label for label, _, _ in MODELS]
               + [f"Số model đúng (/{len(COUNTED)})", "Trong BatDong?", "Trong ToanBoSai?"])
    style_header(ws, headers, [12, 18, 11, 60, 8] + [9] * len(MODELS) + [9, 9, 9])
    ws.freeze_panes = "B2"
    for qid in test_ids:
        row = rows[qid]
        letters = [runs[label].get(qid, {}).get("pred") for label, _, _ in MODELS]
        ws.append([qid, ", ".join(row["medical_topic"]), row["difficulty_level"], row["question"], row["answer"]]
                  + [l or "" for l in letters]
                  + [sum(runs[m][qid]["correct"] for m in COUNTED),
                     "Có" if qid in batdong_set else "", "Có" if qid in toanbosai_set else ""])
        r = ws.max_row
        for cell in ws[r]:
            cell.font = FONT_BODY
        for i, letter in enumerate(letters):
            fill_letter(ws.cell(row=r, column=6 + i), letter, row["answer"])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("LapLuan")
    style_header(ws, ["id", "Danh sách", "Model", "Đáp án chuẩn", "Đáp án lúc trả lời nhanh",
                      "Đáp án khi giải thích", "Câu trả lời có giải thích", "Phần suy nghĩ (thinking)"],
                 [12, 14, 20, 8, 10, 10, 90, 90])
    for qid in batdong + [i for i in toanbosai if i not in batdong_set]:
        where = " + ".join(n for n, s in (("BatDong", batdong_set), ("ToanBoSai", toanbosai_set)) if qid in s)
        for label, recs in reasoning.items():
            rec = recs.get(qid)
            if rec is None:
                continue
            ws.append([qid, where, label, rows[qid]["answer"], runs.get(base_of[label], {}).get(qid, {}).get("pred") or "",
                       rec["pred"] or "Không chọn", clip(rec.get("answer_text")), clip(rec.get("reasoning"))])
            r = ws.max_row
            for cell in ws[r]:
                cell.font = FONT_BODY
                cell.alignment = WRAP_TOP
            fill_letter(ws.cell(row=r, column=6), rec["pred"], rows[qid]["answer"])
    ws.auto_filter.ref = ws.dimensions

    review_ids = batdong + [i for i in toanbosai if i not in batdong_set]
    n_third = sum(third_status(i) == "Có" for i in batdong)
    n_third_answered = sum(third_status(i) != "Chưa trả lời" for i in batdong)
    write_guide(guide, len(test_ids), len(batdong), n_third, n_third_answered, len(toanbosai),
                len(batdong_set & toanbosai_set), len(review_ids), reasoning)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if workbook_path.exists():
        backup = OUT_DIR / "backup" / f"VM14K_review.{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        backup.parent.mkdir(exist_ok=True)
        shutil.copy2(workbook_path, backup)
    wb.save(workbook_path)
    (OUT_DIR / "review_ids.txt").write_text("\n".join(review_ids) + "\n", encoding="utf-8")
    write_markdown(OUT_DIR / "reasoning_log.md", review_ids, rows, runs, reasoning, batdong_set, toanbosai_set)
    print(f"BatDong {len(batdong)} (Nemotron agrees {n_third}/{n_third_answered} answered), "
          f"ToanBoSai {len(toanbosai)}, overlap {len(batdong_set & toanbosai_set)}, "
          f"review ids {len(review_ids)}, explanation logs: "
          + (", ".join(f"{k} {len(v)}" for k, v in reasoning.items()) or "none")
          + f"; reviewer rows kept {sum(len(v) for v in annotations.values())}"
          + (f" ({len(orphans)} moved to ChuThichCu)" if orphans else ""))
    print(f"wrote {os.path.relpath(OUT_DIR, REPO_ROOT)}/VM14K_review.xlsx, review_ids.txt, reasoning_log.md")


def write_guide(ws, n_test, n_bd, n_third, n_third_answered, n_tbs, n_overlap, n_review, reasoning):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    lines = [
        ("VM14K — danh sách câu cần duyệt", ""),
        ("", ""),
        ("Nguồn", f"Tập test cố định (splits/split_v1.json), {n_test} câu, dữ liệu data/cleaned/clean_final.jsonl. "
                  "Đáp án model lấy từ các lần chạy zero-shot trong reports/eval/runs/ "
                  "(prompt của paper, tắt thinking, model chỉ trả lời 1 chữ cái)."),
        ("BatDong", f"{n_bd} câu mà DeepSeek-v4.1-flash và Gemma-4-31B cùng chọn MỘT đáp án, khác đáp án chuẩn. "
                    f"Nemotron-3-Ultra trả lời {n_third_answered} câu trong số này và chọn trùng ở {n_third} câu. "
                    "Sắp xếp: 3 model cùng chọn lên trước. Đây là nhóm nghi đáp án chuẩn sai nhiều nhất."),
        ("ToanBoSai", f"{n_tbs} câu mà cả {len(COUNTED)} model chạy đủ ({', '.join(COUNTED)}) đều sai. "
                      f"{n_overlap} câu nằm ở cả hai danh sách. Sắp xếp theo số model cùng chọn một đáp án sai, nhiều nhất lên trước. "
                      "Khi hai đáp án hòa phiếu, cột 'được chọn nhiều nhất' ghi cả hai, kèm chữ '(hòa)'."),
        ("DoKho", "Độ chính xác theo nhãn độ khó của VM14K. Nếu nhãn có ý nghĩa thì Easy phải cao hơn hẳn Challenging."),
        ("TatCa", "Toàn bộ câu test với đáp án từng model, để lọc tự do (ví dụ: câu Easy mà ≤2 model đúng)."),
        ("LapLuan", "Mỗi dòng là một (câu, model): câu trả lời có giải thích và phần suy nghĩ (thinking) đầy đủ. "
                    "Bản dễ đọc hơn: reports/eval/review/reasoning_log.md."),
        ("", ""),
        ("Màu ô đáp án", "Xanh = trùng đáp án chuẩn, đỏ = khác đáp án chuẩn, xám = model không trả lời được."),
        ("Cột giải thích", "Chạy lại riêng các câu trong 2 danh sách, yêu cầu model giải thích bằng tiếng Việt rồi chốt đáp án "
                           "(scripts/eval/collect_reasoning.py). Prompt khác lần chạy gốc nên model có thể đổi đáp án: "
                           "ô ghi '(đổi từ X)' khi khác lần trả lời nhanh. 'Không chọn' = model nói không có "
                           "phương án đúng, thiếu hình hoặc thiếu dữ kiện. Model có log: "
                           + (", ".join(reasoning) or "chưa có — chạy collect_reasoning.py rồi build lại workbook") + "."),
        ("", ""),
        ("Cách duyệt", "Điền 5 cột vàng ở cuối mỗi dòng. Nên đọc câu hỏi và tự chọn đáp án TRƯỚC khi nhìn đáp án chuẩn "
                       "và đáp án model, để không bị dẫn dắt. Câu nằm ở cả hai sheet: duyệt ở BatDong là đủ."),
        ("Giữ phần đã điền", "Chạy lại script sẽ giữ nguyên 5 cột vàng theo id (đọc từ file cũ) và sao lưu file cũ vào "
                             "reports/eval/review/backup/. Câu đã điền mà không còn trong danh sách được chuyển sang sheet "
                             "ChuThichCu. Đóng file trong Excel trước khi chạy lại, và đừng đổi tên các cột."),
        ("Phân loại", " | ".join(REVIEW_CATEGORIES)),
        ("", ""),
        ("Lưu ý", "• Model bất đồng với đáp án chuẩn chưa có nghĩa đáp án chuẩn sai — đây chỉ là danh sách nghi vấn.\n"
                  "• Phần giải thích/thinking là lời model tự nói, không bảo đảm phản ánh đúng cách model thật sự ra đáp án.\n"
                  "• 62 câu có cờ contradiction_pending_review không nằm trong tập test nên không có ở đây.\n"
                  f"• Tổng số câu cần duyệt (hợp 2 danh sách): {n_review}."),
        ("Tạo lại", "python scripts/eval/build_review_workbook.py"),
    ]
    for key, text in lines:
        ws.append([key, text])
        r = ws.max_row
        ws.cell(row=r, column=1).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2).font = Font(name=FONT)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = TOP
    ws.cell(row=1, column=1).font = Font(name=FONT, bold=True, size=14)


def write_markdown(path, ids, rows, runs, reasoning, batdong_set, toanbosai_set):
    out = ["# VM14K — log giải thích của model", "",
           "Sinh bởi `scripts/eval/build_review_workbook.py`. Mỗi mục: câu hỏi, đáp án chuẩn, đáp án của từng model "
           "lúc trả lời nhanh, rồi phần giải thích khi được yêu cầu giải thích. ✓ = trùng đáp án chuẩn.", ""]
    for n, qid in enumerate(ids, start=1):
        row = rows[qid]
        gold = row["answer"]
        where = " + ".join(n_ for n_, s in (("BatDong", batdong_set), ("ToanBoSai", toanbosai_set)) if qid in s)
        quick = ", ".join(f"{label} {runs[label][qid]['pred'] or '?'}{' ✓' if runs[label][qid]['correct'] else ''}"
                          for label, _, _ in MODELS if qid in runs[label])
        out += [f"## {n}. `{qid}` — {where}", "",
                f"**Chuyên khoa:** {', '.join(row['medical_topic'])} · **Độ khó:** {row['difficulty_level']}", "",
                row["question"], ""]
        out += [f"- **{LETTERS[i]}.** {opt}{'  ← đáp án chuẩn' if LETTERS[i] == gold else ''}"
                for i, opt in enumerate(row["options"])]
        out += ["", f"**Trả lời nhanh:** {quick}", ""]
        for label, recs in reasoning.items():
            rec = recs.get(qid)
            if rec is None:
                continue
            mark = " ✓" if rec["pred"] == gold else ""
            out += [f"### {label} — chốt {rec['pred'] or 'không chọn đáp án nào'}{mark}", "", (rec.get("answer_text") or "").strip(), ""]
            if rec.get("reasoning"):
                out += ["<details><summary>Phần suy nghĩ (thinking)</summary>", "", rec["reasoning"].strip(), "",
                        "</details>", ""]
        out += ["---", ""]
    path.write_text("\n".join(out), encoding="utf-8")


if __name__ == "__main__":
    main()
