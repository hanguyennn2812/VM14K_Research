#!/usr/bin/env python3
"""
build_review_sample.py — draw the questions a clinician reviews, and the blinded
files they review them in.

Two kinds of strata, kept apart because they answer different questions:

    ngau_nhien      simple random sample of the whole test split. The ONLY
                    stratum that estimates the benchmark's answer-key error rate.
    uu_tien_*       priority pools built from model disagreement (see
                    build_review_workbook.py). They find errors efficiently but
                    say nothing about the error rate outside their pool.

Priority pools are disjoint and exclude the random draw; each id goes to the
first pool it matches, in this order:

    uu_tien_1_tu_choi          Nemotron, asked to explain, names no option
    uu_tien_2_dong_thuan_giu   DeepSeek, Gemma-31B and Nemotron pick the same
                               non-key option, and Nemotron keeps it when asked
                               to explain
    uu_tien_3_batdong_khac     rest of BatDong
    uu_tien_4_toanbosai_khac   rest of ToanBoSai

Every row records its pool size and selection probability. About 20% of each
stratum is flagged for a second, independent reviewer.

Review runs in two rounds, in separate files so the first stays blind:

    vong1_doc_lap.xlsx         question + options only, shuffled, under a code
                               (no id, no key, no model answers, no stratum)
    vong1_doc_lap_nguoi2.xlsx  the same, for the second reviewer's rows only
    vong2_doi_chieu.xlsx       opened only after round 1: key, model answers,
                               Nemotron's explanation, adjudication fields
    sample_manifest.csv        code ↔ id ↔ stratum ↔ probability. Not for reviewers.

The draw is frozen: the script refuses to overwrite existing files (they hold
reviewer input) unless --force.

    python scripts/eval/build_review_sample.py
"""
from __future__ import annotations

import argparse
import collections
import csv
import math
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_workbook import (  # noqa: E402
    COUNTED, FILL_REVIEW, FONT, FONT_BODY, MODELS, REPO_ROOT, STRONG_PAIR, THIRD, WRAP_TOP, clip,
    fill_letter, load_reasoning, load_rows, load_run, option_text, options_block, style_header,
)

OUT_DIR = REPO_ROOT / "reports" / "eval" / "review_sample"
EXPLAINER = "Nemotron-3-Ultra"  # the model whose explanation log exists for every review id

ANSWER_CHOICES = ["A", "B", "C", "D", "E", "F", "G",
                  "Nhiều đáp án đúng", "Không có đáp án đúng", "Không xác định được"]
# One exclusive overall status (what endpoints and kappa use) plus independent defect flags,
# because one question can have several defects at once.
STATUS_CHOICES = ["Bình thường", "Có lỗi nhưng vẫn trả lời được", "Không trả lời được"]
DEFECT_FLAGS = [("Lỗi: mơ hồ", 9), ("Lỗi: thiếu dữ kiện", 9), ("Lỗi: thiếu hình ảnh", 9),
                ("Lỗi: chính tả - đánh máy", 9), ("Lỗi: phương án lỗi hoặc trùng", 9)]
CAUSES = ["Đáp án chuẩn sai", "Câu hỏi lỗi", "Model thiếu kiến thức VN", "Model sai kiến thức chung",
          "Không rõ", "Không áp dụng"]

ROUND1_FIELDS = [
    ("Đáp án bạn chọn", 14, ANSWER_CHOICES),
    ("Tập đáp án đúng (nếu nhiều, ví dụ AC)", 12, None),
    ("Mức chắc chắn", 12, ["Chắc chắn", "Khá chắc", "Không chắc"]),
    ("Tình trạng câu hỏi", 16, STATUS_CHOICES),
    *[(name, w, ["Có"]) for name, w in DEFECT_FLAGS],
    ("Phụ thuộc phác đồ / thực hành VN?", 12, ["Có", "Không", "Không rõ"]),
    ("Nguồn tham khảo", 30, None),
    ("Ghi chú", 30, None),
    ("Người duyệt", 12, None),
]
ROUND2_FIELDS = [
    ("Trạng thái đáp án chuẩn", 18, ["Đúng", "Sai", "Là một trong nhiều đáp án đúng", "Không xác định"]),
    ("Đáp án đúng cuối cùng", 14, ANSWER_CHOICES),
    ("Tập đáp án đúng cuối cùng (nếu nhiều)", 12, None),
    ("Tình trạng câu hỏi cuối cùng", 16, STATUS_CHOICES),
    *[(name, w, ["Có"]) for name, w in DEFECT_FLAGS],
    ("Bối cảnh kiến thức", 18, ["Chuẩn chung quốc tế", "Theo phác đồ - thực hành VN", "Không rõ", "Không áp dụng"]),
    ("Kiến thức đã thay đổi theo thời gian?", 12, ["Có", "Không", "Không rõ"]),
    ("Vì sao đáp án đa số model khác đáp án chuẩn — chính", 20, CAUSES),
    ("— phụ (nếu có)", 16, CAUSES),
    ("Bằng chứng / nguồn", 30, None),
    ("Ghi chú", 30, None),
    ("Người duyệt", 12, None),
]


def add_fields(ws, fields, first_col, n_rows):
    """Reviewer columns: shaded, wrapped, with a dropdown where the field has fixed choices."""
    for k, (_, _, choices) in enumerate(fields):
        col = first_col + k
        if choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{n_rows + 1}")
        for r in range(2, n_rows + 2):
            cell = ws.cell(row=r, column=col)
            cell.fill = FILL_REVIEW
            cell.alignment = WRAP_TOP


def write_round1(path, sample, rows, title):
    wb = Workbook()
    guide = wb.active
    guide.title = "HuongDan"
    guide.column_dimensions["A"].width = 120
    for text in [
        title,
        "",
        "Vòng 1 — duyệt độc lập. Với mỗi câu, hãy tự trả lời như khi làm đề: chọn đáp án đúng theo kiến thức "
        "và tài liệu của bạn, ghi mức chắc chắn và tình trạng câu hỏi.",
        "File này cố ý KHÔNG có đáp án chuẩn của bộ đề, đáp án của các model, hay lý do câu được chọn. "
        "Đừng tra cứu câu hỏi trong các file khác của dự án trước khi xong vòng này.",
        "Nếu có nhiều đáp án đúng: chọn 'Nhiều đáp án đúng' và ghi các chữ vào cột kế bên (ví dụ AC). "
        "Nếu không đủ dữ kiện để trả lời: chọn 'Không xác định được'.",
        "'Tình trạng câu hỏi' chọn MỘT mức tổng quát. Các cột 'Lỗi: …' đánh 'Có' cho MỌI loại lỗi gặp phải "
        "(một câu có thể có nhiều lỗi); khi đã chọn 'Tình trạng câu hỏi', ô cờ để trống nghĩa là KHÔNG có lỗi đó. Lỗi đánh máy không ảnh hưởng nghĩa → "
        "'Có lỗi nhưng vẫn trả lời được'.",
        "Được phép tra sách, phác đồ Bộ Y tế, hướng dẫn chuyên ngành — ghi nguồn vào 'Nguồn tham khảo'.",
        "Xong vòng 1 thì gửi lại file và KHÔNG sửa nó nữa. Vòng 2 chỉ được gửi khi CẢ HAI người duyệt đã nộp "
        "vòng 1. Không xem file vòng 1 của người duyệt kia.",
    ]:
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row=guide.max_row, column=1).font = Font(name=FONT)
    guide.cell(row=1, column=1).font = Font(name=FONT, bold=True, size=14)

    ws = wb.create_sheet("Vong1")
    base = [("Mã câu", 8), ("Chuyên khoa", 18), ("Câu hỏi", 55), ("Các phương án", 60)]
    style_header(ws, [h for h, _ in base] + [h for h, _, _ in ROUND1_FIELDS],
                 [w for _, w in base] + [w for _, w, _ in ROUND1_FIELDS])
    for s in sample:
        row = rows[s["id"]]
        ws.append([s["code"], ", ".join(row["medical_topic"]), row["question"], options_block(row)]
                  + [""] * len(ROUND1_FIELDS))
        for cell in ws[ws.max_row]:
            cell.font = FONT_BODY
            cell.alignment = WRAP_TOP
    add_fields(ws, ROUND1_FIELDS, len(base) + 1, len(sample))
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_round2(path, sample, rows, runs, explain, majority, title):
    wb = Workbook()
    guide = wb.active
    guide.title = "HuongDan"
    guide.column_dimensions["A"].width = 120
    for text in [
        title,
        "",
        "Chỉ mở SAU khi cả hai người duyệt đã nộp vòng 1. Làm độc lập, không trao đổi với người duyệt kia.",
        "Mỗi câu giờ có thêm đáp án chuẩn của bộ đề, đáp án của các model (trả lời nhanh, tắt thinking), đáp án "
        f"đa số của {len(COUNTED)} model, và phần giải thích của {EXPLAINER} khi được yêu cầu lập luận. "
        "Xanh = trùng đáp án chuẩn, đỏ = khác.",
        "Mở file vòng 1 của bạn bên cạnh (cùng 'Mã câu'). Nếu đổi ý so với vòng 1, KHÔNG sửa file vòng 1 — "
        "ghi lý do vào cột 'Ghi chú' ở đây.",
        "Các trường tách riêng, không loại trừ nhau: một câu có thể vừa lỗi đề vừa sai đáp án, vừa theo phác đồ "
        "VN vừa là kiến thức đã thay đổi. Cột 'Vì sao…' nói về ĐÁP ÁN ĐA SỐ của model (cột riêng); nếu đa số "
        "model trùng đáp án chuẩn thì chọn 'Không áp dụng'.",
        "Phần giải thích của model là lời model tự nói, có thể sai — dùng để gợi ý, không phải bằng chứng.",
    ]:
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row=guide.max_row, column=1).font = Font(name=FONT)
    guide.cell(row=1, column=1).font = Font(name=FONT, bold=True, size=14)

    ws = wb.create_sheet("Vong2")
    base = ([("Mã câu", 8), ("Chuyên khoa", 18), ("Câu hỏi", 50), ("Các phương án", 55),
             ("Đáp án chuẩn", 8), ("Nội dung đáp án chuẩn", 28)]
            + [(label, 9) for label, _, _ in MODELS]
            + [(f"Đáp án đa số {len(COUNTED)} model", 11)]
            + [(f"{EXPLAINER} khi giải thích", 11), (f"{EXPLAINER} — giải thích", 70)])
    style_header(ws, [h for h, _ in base] + [h for h, _, _ in ROUND2_FIELDS],
                 [w for _, w in base] + [w for _, w, _ in ROUND2_FIELDS])
    model_col0 = 7
    for s in sample:
        row = rows[s["id"]]
        gold = row["answer"]
        letters = [runs[label].get(s["id"], {}).get("pred") for label, _, _ in MODELS]
        top, votes = majority(s["id"])
        rec = explain.get(s["id"])
        ws.append([s["code"], ", ".join(row["medical_topic"]), row["question"], options_block(row),
                   gold, option_text(row, gold)]
                  + [l or "" for l in letters]
                  + [" / ".join(top) + f" ({votes}/{len(COUNTED)}{', hòa' if len(top) > 1 else ''})"]
                  + ([rec["pred"] or "Không chọn", clip(rec.get("answer_text"))] if rec else ["", ""])
                  + [""] * len(ROUND2_FIELDS))
        r = ws.max_row
        for cell in ws[r]:
            cell.font = FONT_BODY
            cell.alignment = WRAP_TOP
        for i, letter in enumerate(letters):
            fill_letter(ws.cell(row=r, column=model_col0 + i), letter, gold)
        fill_letter(ws.cell(row=r, column=model_col0 + len(MODELS)), top[0] if len(top) == 1 else None, gold)
        if rec:
            fill_letter(ws.cell(row=r, column=model_col0 + len(MODELS) + 1), rec["pred"], gold)
            ws.cell(row=r, column=model_col0 + len(MODELS) + 2).alignment = Alignment(vertical="top")
    add_fields(ws, ROUND2_FIELDS, len(base) + 1, len(sample))
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--n-random", type=int, default=60)
    ap.add_argument("--n-batdong-khac", type=int, default=15)
    ap.add_argument("--n-toanbosai-khac", type=int, default=15)
    ap.add_argument("--double-share", type=float, default=0.2, help="share of each stratum for a 2nd reviewer")
    ap.add_argument("--force", action="store_true", help="overwrite existing files (they may hold reviewer input)")
    args = ap.parse_args()

    outputs = [OUT_DIR / n for n in ("sample_manifest.csv", "vong1_doc_lap.xlsx", "vong1_doc_lap_nguoi2.xlsx",
                                     "vong2_doi_chieu.xlsx", "vong2_doi_chieu_nguoi2.xlsx")]
    existing = [p for p in outputs if p.exists()]
    if existing and not args.force:
        sys.exit("refusing to overwrite (may hold reviewer input): "
                 + ", ".join(str(p.relative_to(REPO_ROOT)) for p in existing) + " — pass --force")

    rows = load_rows()
    runs = {label: load_run(stem) for label, stem, _ in MODELS}
    reasoning, _ = load_reasoning()
    if EXPLAINER not in reasoning:
        sys.exit(f"no {EXPLAINER} explanation log in reports/eval/reasoning/ — run collect_reasoning.py first")
    explain = reasoning[EXPLAINER]
    test_ids = sorted(runs[STRONG_PAIR[0]])

    a, b = (runs[m] for m in STRONG_PAIR)
    third = runs[THIRD]
    batdong = {i for i in test_ids if i in b and a[i]["pred"] and a[i]["pred"] == b[i]["pred"] != a[i]["gold"]}
    toanbosai = {i for i in test_ids if not any(runs[m][i]["correct"] for m in COUNTED)}
    # Defined on the review list only, so explanations collected later for random-sample questions
    # cannot change the priority pools.
    abstained = {i for i, rec in explain.items()
                 if rec.get("status") == "abstained" and i in batdong | toanbosai}
    kept = {i for i in batdong if i in third and third[i]["pred"] == a[i]["pred"]
            and i in explain and explain[i]["pred"] == a[i]["pred"]}

    rng = random.Random(args.seed)
    sample = []

    def draw(stratum, pool, n):
        pool = sorted(pool)
        picked = pool if n >= len(pool) else rng.sample(pool, n)
        for i in picked:
            sample.append({"id": i, "stratum": stratum, "pool_size": len(pool), "n_drawn": len(picked),
                           "stage_selection_prob": round(len(picked) / len(pool), 6) if pool else 0.0})
        return set(picked)

    priority_sets = [
        ("uu_tien_1_tu_choi", abstained, None),
        ("uu_tien_2_dong_thuan_giu", kept, None),
        ("uu_tien_3_batdong_khac", batdong, args.n_batdong_khac),
        ("uu_tien_4_toanbosai_khac", toanbosai, args.n_toanbosai_khac),
    ]

    def priority_set_of(qid):
        """The priority set a question belongs to by membership (first match), drawn or not."""
        return next((name for name, members, _ in priority_sets if qid in members), "")

    random_ids = draw("ngau_nhien", test_ids, args.n_random)
    # Pools are disjoint by membership (first match wins), not by what happened to be drawn,
    # and none contains a question already in the random sample.
    claimed = set(random_ids)
    for stratum, members, n in priority_sets:
        pool = set(members) - claimed
        draw(stratum, pool, len(pool) if n is None else n)
        claimed |= set(members)
    for s in sample:
        s["priority_set"] = priority_set_of(s["id"])

    def majority(qid):
        votes = collections.Counter(runs[m][qid]["pred"] for m in COUNTED if runs[m][qid]["pred"])
        if not votes:
            return [], 0
        top = max(votes.values())
        return sorted(letter for letter, n in votes.items() if n == top), top

    # Second reviewer: ceil(share) of every stratum, drawn independently per stratum.
    by_stratum = {}
    for s in sample:
        by_stratum.setdefault(s["stratum"], []).append(s)
    for members in by_stratum.values():
        k = math.ceil(args.double_share * len(members))
        flagged = {s["id"] for s in rng.sample(members, k)}
        for s in members:
            s["duyet_doi"] = int(s["id"] in flagged)

    # Reviewers see questions in random order under a code, so neither order nor id reveals the stratum.
    rng.shuffle(sample)
    for n, s in enumerate(sample, start=1):
        s["code"] = f"Q{n:03d}"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with open(OUT_DIR / "sample_manifest.csv", "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "id", "stratum", "pool_size", "n_drawn", "stage_selection_prob",
                    "priority_set", "duyet_doi", "seed"])
        for s in sample:
            w.writerow([s["code"], s["id"], s["stratum"], s["pool_size"], s["n_drawn"], s["stage_selection_prob"],
                        s["priority_set"], s["duyet_doi"], args.seed])
    second = [s for s in sample if s["duyet_doi"]]
    write_round1(OUT_DIR / "vong1_doc_lap.xlsx", sample, rows, "Vòng 1 — người duyệt chính")
    write_round1(OUT_DIR / "vong1_doc_lap_nguoi2.xlsx", second, rows,
                 "Vòng 1 — người duyệt thứ hai (làm độc lập, không trao đổi với người duyệt chính)")
    write_round2(OUT_DIR / "vong2_doi_chieu.xlsx", sample, rows, runs, explain, majority,
                 "Vòng 2 — đối chiếu, người duyệt chính")
    write_round2(OUT_DIR / "vong2_doi_chieu_nguoi2.xlsx", second, rows, runs, explain, majority,
                 "Vòng 2 — đối chiếu, người duyệt thứ hai")

    # Round 2 should give every question the same information, so explanations missing for sampled
    # questions (the random stratum reaches outside the review list) are listed for collect_reasoning.py.
    no_explain = [s["id"] for s in sample if s["id"] not in explain]
    missing_path = OUT_DIR / "explain_missing_ids.txt"
    if no_explain:
        missing_path.write_text("\n".join(no_explain) + "\n", encoding="utf-8")
    elif missing_path.exists():
        missing_path.unlink()

    print(f"{len(sample)} questions, {len(second)} double-reviewed:")
    for stratum, members in by_stratum.items():
        m = members[0]
        print(f"  {stratum}: {len(members)} of pool {m['pool_size']} (stage p={m['stage_selection_prob']}), "
              f"{sum(s['duyet_doi'] for s in members)} double")
    if no_explain:
        print(f"warning: {len(no_explain)} sampled questions have no {EXPLAINER} explanation — run "
              f"collect_reasoning.py --ids {missing_path.relative_to(REPO_ROOT)} --tag review_ids, then --force")
    print(f"wrote {OUT_DIR.relative_to(REPO_ROOT)}/")


if __name__ == "__main__":
    main()
