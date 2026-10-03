#!/usr/bin/env python3
"""Phân tích kế hoạch đã chốt trong HUONG_DAN.md; không sửa đầu vào.

python scripts/eval/analyze_review.py [--dir PATH] [--simulate] [--seed 42]
Phân xử: sheet PhanXu, Mã câu + các trường ROUND2_FIELDS[:4] + Nguồn +
Người phân xử. Sao chép phan_xu_template.xlsx thành phan_xu.xlsx để điền.
"""
from __future__ import annotations

import argparse
from collections import Counter
import math
from pathlib import Path
import re
import shutil
import sys
import tempfile

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_sample import (  # noqa: E402
    ROUND1_FIELDS, ROUND2_FIELDS, ANSWER_CHOICES, STATUS_CHOICES, DEFECT_FLAGS, OUT_DIR,
)

R1_ANSWER, R1_SET = [f[0] for f in ROUND1_FIELDS[:2]]
R1_STATUS = ROUND1_FIELDS[3][0]
KEY_STATUS, R2_ANSWER, R2_SET, R2_STATUS = [f[0] for f in ROUND2_FIELDS[:4]]
MULTI, NONE, UNKNOWN = ANSWER_CHOICES[-3:]
KEY_OK, KEY_WRONG, KEY_MULTI, KEY_UNKNOWN = ROUND2_FIELDS[0][2]
CODE, OPTIONS = "Mã câu", "Các phương án"
SOURCE, ADJUDICATOR = "Nguồn", "Người phân xử"
FILES = (
    ("vong1_doc_lap.xlsx", "Vong1", 1, 1),
    ("vong1_doc_lap_nguoi2.xlsx", "Vong1", 1, 2),
    ("vong2_doi_chieu.xlsx", "Vong2", 2, 1),
    ("vong2_doi_chieu_nguoi2.xlsx", "Vong2", 2, 2),
)


def clean(value):
    return "" if value is None or pd.isna(value) else str(value).strip()


def normalize_answer(choice, letters="", n_options=None):
    """Return a hashable sorted tuple, a distinct special category, or None.

    Reject malformed combinations rather than turning them into missing answers.
    Separators in sets may be whitespace, comma, slash, or semicolon.
    """
    choice, letters = clean(choice), clean(letters)
    if letters and choice != MULTI:
        raise ValueError("có tập chữ nhưng không chọn Nhiều đáp án đúng")
    if not choice:
        return None
    if choice not in ANSWER_CHOICES:
        raise ValueError(f"giá trị đáp án ngoài dropdown: {choice}")
    if choice == MULTI:
        if not letters:
            raise ValueError("Nhiều đáp án đúng nhưng thiếu tập chữ")
        compact = re.sub(r"[\s,;/]+", "", letters.upper())
        if not re.fullmatch(r"[A-Z]+", compact):
            raise ValueError(f"tập chữ không hợp lệ: {letters}")
        answer = tuple(sorted(set(compact)))
        if len(answer) < 2:
            raise ValueError("Nhiều đáp án đúng cần ít nhất hai chữ khác nhau")
    elif choice in (NONE, UNKNOWN):
        return choice
    else:
        answer = (choice,)
    if n_options is not None and any(ord(x) - ord("A") >= n_options for x in answer):
        raise ValueError(f"chữ vượt quá {n_options} phương án: {''.join(answer)}")
    return answer


def validate_row(row, fields, n_options):
    errors = []
    for name, _, choices in fields:
        value = clean(row.get(name))
        if value and choices and value not in choices:
            errors.append(f"{name}: giá trị ngoài dropdown ({value})")
    answer_field, set_field = (R1_ANSWER, R1_SET) if fields is ROUND1_FIELDS else (R2_ANSWER, R2_SET)
    try:
        answer = normalize_answer(row.get(answer_field), row.get(set_field), n_options)
    except ValueError as exc:
        errors.append(str(exc))
        answer = None
    return answer, errors


def flag_value(row, flag, status_field):
    """Defect flags are only ever marked "Có". Once the overall status is filled, a blank flag
    means that defect is absent ("Không"); with no status, the flag is missing (HUONG_DAN.md §3)."""
    if clean(row.get(flag)) == "Có":
        return "Có"
    return "Không" if clean(row.get(status_field)) else None


def flag_consistency(row, status_field):
    """Status and flags that contradict each other. Reported, but the row stays usable."""
    status = clean(row.get(status_field))
    marked = [name for name, _ in DEFECT_FLAGS if clean(row.get(name)) == "Có"]
    if status == STATUS_CHOICES[0] and marked:
        return [f"cảnh báo: tình trạng 'Bình thường' nhưng có cờ lỗi ({', '.join(marked)})"]
    if status in STATUS_CHOICES[1:] and not marked:
        return ["cảnh báo: tình trạng có lỗi nhưng không đánh cờ loại lỗi nào"]
    if marked and not status:
        return ["cảnh báo: có cờ lỗi nhưng thiếu tình trạng tổng quát"]
    return []


def wilson(successes, n):
    if not n:
        return (math.nan, math.nan)
    z = 1.959963984540054
    p = successes / n
    denominator = 1 + z * z / n
    center = (p + z * z / (2 * n)) / denominator
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denominator
    return max(0., center - half), min(1., center + half)


def cohen_kappa(left, right):
    if len(left) != len(right):
        raise ValueError("paired lengths differ")
    n = len(left)
    if not n:
        return math.nan
    a, b = Counter(left), Counter(right)
    observed = sum(x == y for x, y in zip(left, right)) / n
    expected = sum(a[x] * b[x] for x in a.keys() | b.keys()) / n ** 2
    # A single constant category has raw agreement 1, but undefined kappa.
    return (observed - expected) / (1 - expected) if expected < 1 else math.nan


def agreement(pairs, seed=42, bootstraps=2000):
    complete = [(a, b) for a, b in pairs if a is not None and b is not None]
    n = len(complete)
    result = {"n": n, "missing": len(pairs) - n, "raw": math.nan,
              "kappa": math.nan, "raw_ci": (math.nan, math.nan),
              "kappa_ci": (math.nan, math.nan), "undefined_bootstraps": 0}
    if not n:
        return result
    left, right = zip(*complete)
    result.update(raw=sum(a == b for a, b in complete) / n, kappa=cohen_kappa(left, right))
    rng = np.random.default_rng(seed)
    raws, kappas = [], []
    for _ in range(bootstraps):
        indices = rng.integers(0, n, n)
        a, b = [left[i] for i in indices], [right[i] for i in indices]
        raws.append(sum(x == y for x, y in zip(a, b)) / n)
        k = cohen_kappa(a, b)
        if math.isfinite(k):
            kappas.append(k)
    result["raw_ci"] = tuple(np.quantile(raws, [.025, .975]))
    result["undefined_bootstraps"] = bootstraps - len(kappas)
    if kappas:
        result["kappa_ci"] = tuple(np.quantile(kappas, [.025, .975]))
    return result


def read_table(path, sheet, required, problems):
    """Keep duplicate rows out of analysis, but report every affected code."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(wb[sheet].iter_rows(values_only=True))
        finally:
            wb.close()
    except (OSError, KeyError, ValueError) as exc:
        problems.append(("*", path.name, f"không đọc được: {exc}"))
        return {}
    if not rows:
        problems.append(("*", path.name, "sheet rỗng"))
        return {}
    headers = [clean(x) for x in rows[0]]
    absent = set(required) - set(headers)
    duplicates = [x for x, count in Counter(headers).items() if x and count > 1]
    if absent or duplicates:
        problems.append(("*", path.name, f"thiếu cột {sorted(absent)}; cột trùng {duplicates}"))
        return {}
    grouped = {}
    for number, values in enumerate(rows[1:], 2):
        if not any(clean(x) for x in values):
            continue
        row = dict(zip(headers, values))
        code = clean(row.get(CODE))
        if not code:
            problems.append(("*", path.name, f"dòng {number}: thiếu mã câu"))
            continue
        grouped.setdefault(code, []).append(row)
    result = {}
    for code, values in grouped.items():
        if len(values) != 1:
            problems.append((code, path.name, f"mã trùng: {len(values)} dòng"))
        else:
            result[code] = values[0]
    return result


def load_inputs(directory):
    problems = []
    manifest = pd.read_csv(directory / "sample_manifest.csv", dtype=str, keep_default_na=False)
    needed = {"code", "stratum", "pool_size", "priority_set", "duyet_doi"}
    if not needed.issubset(manifest.columns):
        raise ValueError(f"manifest thiếu cột: {sorted(needed - set(manifest.columns))}")
    manifest["code"] = manifest["code"].str.strip()
    if manifest["code"].eq("").any() or manifest["code"].duplicated().any():
        bad = manifest.loc[manifest["code"].eq("") | manifest["code"].duplicated(False), "code"].tolist()
        raise ValueError(f"manifest mã thiếu/trùng (không thể xác định mẫu số): {bad}")
    if not manifest["duyet_doi"].isin(["0", "1"]).all():
        raise ValueError("manifest duyet_doi phải là 0 hoặc 1")
    records = {}
    for filename, sheet, round_, reviewer in FILES:
        fields = ROUND1_FIELDS if round_ == 1 else ROUND2_FIELDS
        expected = set(manifest.loc[(manifest.duyet_doi == "1") | (reviewer == 1), "code"])
        table = read_table(directory / filename, sheet, [CODE, OPTIONS] + [f[0] for f in fields], problems)
        for code in sorted(set(table) - expected):
            problems.append((code, filename, "mã ngoài danh sách dự kiến"))
        for code in sorted(expected):
            row = table.get(code)
            if row is None:
                problems.append((code, filename, "thiếu dòng (hoặc dòng trùng không dùng được)"))
                records[round_, reviewer, code] = {}
                continue
            n_options = len(clean(row.get(OPTIONS)).splitlines()) if clean(row.get(OPTIONS)) else 0
            answer, errors = validate_row(row, fields, n_options)
            if n_options == 0:
                errors.append("thiếu các phương án")
            for error in errors:
                problems.append((code, filename, error))
            for warning in flag_consistency(row, R1_STATUS if round_ == 1 else R2_STATUS):
                problems.append((code, filename, warning))
            records[round_, reviewer, code] = dict(row, answer=answer, valid=not errors,
                filled=any(clean(row.get(f[0])) for f in fields), n_options=n_options)
    return manifest, records, problems


def label_from(row):
    if not row.get("valid"):
        return {"key": None, "answer": None, "status": None}
    return {"key": clean(row.get(KEY_STATUS)) or None, "answer": row.get("answer"),
            "status": clean(row.get(R2_STATUS)) or None}


def final_labels(directory, manifest, records, problems):
    decisions = {}
    if (directory / "phan_xu.xlsx").exists():
        decisions = read_table(directory / "phan_xu.xlsx", "PhanXu",
            [CODE] + [f[0] for f in ROUND2_FIELDS[:4]] + [SOURCE, ADJUDICATOR], problems)
    final, disagreements = {}, []
    for item in manifest.to_dict("records"):
        code = item["code"]
        first = records.get((2, 1, code), {})
        a = label_from(first)
        flag = "1 người duyệt"
        if item["duyet_doi"] == "1":
            b = label_from(records.get((2, 2, code), {}))
            complete = all(x["key"] is not None and x["answer"] is not None for x in (a, b))
            # A disagreement requires two observed labels; an absent review is missing. A status-only
            # disagreement is adjudicated too, since E2 depends on the final status.
            disagree = any(a[f] is not None and b[f] is not None and a[f] != b[f]
                           for f in ("key", "answer", "status"))
            if disagree:
                disagreements.append((code, a.copy(), b.copy()))
                a = {"key": None, "answer": None, "status": None}
                flag = "chưa phân xử"
                decision = decisions.get(code)
                if decision and any(clean(decision.get(f[0])) for f in ROUND2_FIELDS[:4]):
                    answer, errors = validate_row(decision, ROUND2_FIELDS, first.get("n_options", 0))
                    if not all(clean(decision.get(f)) for f in (KEY_STATUS, R2_ANSWER, R2_STATUS, SOURCE, ADJUDICATOR)):
                        errors.append("quyết định thiếu trạng thái/đáp án/tình trạng/nguồn/người phân xử")
                    for error in errors:
                        problems.append((code, "phan_xu.xlsx", error))
                    if not errors:
                        a = {"key": clean(decision[KEY_STATUS]), "answer": answer,
                             "status": clean(decision[R2_STATUS])}
                        flag = "đã phân xử"
            elif complete:
                flag = "2 người đồng ý"
                if a["status"] is None:  # only one reviewer gave a status: use it
                    a["status"] = b["status"]
            else:
                a = {"key": None, "answer": None, "status": None}
                flag = "thiếu duyệt đôi"
        final[code] = dict(a, flag=flag)
    disagreement_codes = {x[0] for x in disagreements}
    for code in decisions.keys() - disagreement_codes:
        problems.append((code, "phan_xu.xlsx", "quyết định không thuộc danh sách bất đồng hiện tại; không dùng"))
    write_adjudication_template(directory, disagreements)
    return final, disagreements


def show_answer(answer):
    return "∅ (thiếu)" if answer is None else "".join(answer) if isinstance(answer, tuple) else answer


def write_adjudication_template(directory, disagreements):
    wb = Workbook()
    ws = wb.active
    ws.title = "PhanXu"
    ws.append([CODE, "Khoá người 1", "Đáp án người 1", "Tình trạng người 1",
               "Khoá người 2", "Đáp án người 2", "Tình trạng người 2"]
              + [f[0] for f in ROUND2_FIELDS[:4]] + [SOURCE, ADJUDICATOR, "Lý do"])
    for code, a, b in disagreements:
        ws.append([code, a["key"], show_answer(a["answer"]), a["status"],
                   b["key"], show_answer(b["answer"]), b["status"]] + [None] * 7)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(directory / "phan_xu_template.xlsx")
    wb.close()


def endpoints(label):
    key, status = label["key"], label["status"]
    e1 = None if key is None else key == KEY_WRONG
    if status == STATUS_CHOICES[-1] or key == KEY_MULTI:
        e2 = True
    elif status is not None and key is not None:
        e2 = False
    else:
        e2 = None
    e3 = True if e1 is True or e2 is True else False if e1 is False and e2 is False else None
    return e1, e2, e3


def summarize_endpoints(codes, final):
    labels = [final[c] for c in codes]
    values = [endpoints(x) for x in labels]
    n = len(labels)
    counts = [sum(x[i] is True for x in values) for i in range(3)]
    missing = [sum(x[i] is None for x in values) for i in range(3)]
    u = sum(x["key"] == KEY_UNKNOWN for x in labels)
    return {"n": n, "counts": counts, "missing": missing, "u": u,
            "ci": [wilson(s, n) for s in counts],
            "sensitivity": (counts[0] / n, (counts[0] + u) / n) if n else (math.nan, math.nan)}


def majority_reference(row):
    raw = clean(row.get("Đáp án đa số 7 model"))
    prefix = raw.split("(", 1)[0].strip()
    if not re.fullmatch(r"[A-Z](?:\s*/\s*[A-Z])*", prefix):
        return None
    return tuple(sorted(set(re.findall(r"[A-Z]", prefix))))


def change_direction(before, after, reference):
    if reference is None:
        return "thiếu mốc"
    # Membership preserves tied pluralities without calling them one winner.
    def matches(answer):
        return isinstance(answer, tuple) and bool(set(answer) & set(reference))
    a, b = matches(before), matches(after)
    return "về phía mốc" if not a and b else "rời mốc" if a and not b else "giữ quan hệ với mốc"


def pct(x):
    return "không xác định" if not math.isfinite(x) else f"{100*x:.1f}%"


def ci_text(interval, percent=True):
    formatter = pct if percent else lambda x: f"{x:.3f}" if math.isfinite(x) else "không xác định"
    return f"[{formatter(interval[0])}; {formatter(interval[1])}]"


def safe(value):
    return clean(value).replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def analyze(directory, seed=42, bootstraps=2000):
    directory = Path(directory)
    manifest, records, problems = load_inputs(directory)
    final, disagreements = final_labels(directory, manifest, records, problems)
    lines = ["# Kết quả duyệt VM14K", "", "Phạm vi: 1.658 câu test; kế hoạch HUONG_DAN.md mục 1 và 4.", ""]
    any_filled = any(r.get("filled") for r in records.values())
    if not any_filled:
        lines += ["**chưa có dữ liệu duyệt**", ""]
    lines += ["| Vòng / người | Dự kiến | Có điền ít nhất một ô | Đáp án hợp lệ | Chưa điền | Dòng thiếu/lỗi |",
              "|---|---:|---:|---:|---:|---:|"]
    for _, _, round_, reviewer in FILES:
        rows = [r for (v, p, _), r in records.items() if (v, p) == (round_, reviewer)]
        filled = sum(bool(r.get("filled")) for r in rows)
        valid = sum(r.get("valid", False) and r.get("answer") is not None for r in rows)
        invalid = sum(not r.get("valid", False) for r in rows)
        lines.append(f"| {round_} / {reviewer} | {len(rows)} | {filled} | {valid} | {len(rows)-filled} | {invalid} |")
    lines += ["", f"Nhãn cuối: {sum(x['key'] is not None and x['answer'] is not None and x['status'] is not None for x in final.values())}/{len(final)} câu đủ ba trường; "
              f"{len(disagreements)} bất đồng, {sum(x['flag'] == 'chưa phân xử' for x in final.values())} chưa phân xử.",
              "", "## Endpoint", "",
              "Mẫu số là toàn bộ số câu đã bốc trong tầng. Số dương tính dưới đây là đã xác nhận; khi còn thiếu, "
              "tỉ lệ và Wilson chỉ là kết quả tạm thời (không coi ô thiếu là âm tính). "
              "Khoảng nhạy E1 chỉ cộng trạng thái Không xác định; thiếu/chưa phân xử báo riêng, không gộp vào u. "
              "Nếu cả tầng chưa có nhãn endpoint, không xuất tỉ lệ hoặc CI như một kết quả duyệt.", "",
              "| Tầng | n | Endpoint | s/n | Wilson 95% | Còn thiếu endpoint |",
              "|---|---:|---|---|---|---:|"]
    summaries = {}
    # The random stratum carries the error-rate estimate, so it is reported first.
    strata = sorted(manifest.stratum.unique(), key=lambda s: (s != "ngau_nhien", list(manifest.stratum).index(s)))
    for stratum in strata:
        group = manifest[manifest.stratum == stratum]
        summary = summarize_endpoints(group.code.tolist(), final)
        summaries[stratum] = summary
        for i, name in enumerate(("E1", "E2", "E3")):
            s, n = summary["counts"][i], summary["n"]
            no_data = summary["missing"][i] == n
            rate_text = "chưa có dữ liệu" if no_data else f"{s}/{n} = {pct(s/n)}"
            interval = "chưa có dữ liệu" if no_data else ci_text(summary['ci'][i])
            lines.append(f"| {stratum} | {n} | {name} | {rate_text} | {interval} | {summary['missing'][i]} |")
        sensitivity = "chưa có dữ liệu" if summary['missing'][0] == summary['n'] else ci_text(summary['sensitivity'])
        lines += ["", f"{stratum}: E1 u={summary['u']}; khoảng nhạy s/n…(s+u)/n = {sensitivity}.", ""]
        if any(summary["missing"]):
            lines += ["Giới hạn mô tả khi lần lượt coi mọi endpoint còn thiếu là âm/dương (không phải CI): "
                      + "; ".join(f"E{i+1} {ci_text((summary['counts'][i]/summary['n'], (summary['counts'][i]+summary['missing'][i])/summary['n']))}"
                                  for i in range(3)), ""]
        # Keep Markdown tables intact after the per-stratum sensitivity paragraph.
        lines += ["| Tầng | n | Endpoint | s/n | Wilson 95% | Còn thiếu endpoint |",
                  "|---|---:|---|---|---|---:|"]
    lines = lines[:-2]  # final unused table header
    lines += ["## Ước lượng lỗi trong tập ưu tiên S", "",
              "Dùng E3 (hợp khoá sai/đề lỗi); báo E1 và E2 cùng công thức để minh bạch. "
              "Không dùng stage_selection_prob làm trọng số cho toàn mẫu. Khi còn thiếu, đây là số lỗi xác nhận tạm thời.", "",
              "| S | Endpoint | Lỗi ở thành viên ngẫu nhiên | Pool × s/n ưu tiên | lỗi(S) | Còn thiếu (ngẫu nhiên / ưu tiên) |",
              "|---|---|---:|---|---:|---|"]
    for stratum, group in manifest[manifest.stratum != "ngau_nhien"].groupby("stratum", sort=False):
        pools = set(group.pool_size)
        pool = int(next(iter(pools))) if len(pools) == 1 else None
        if pool is None:
            problems.append(("*", "manifest", f"{stratum}: pool_size không nhất quán"))
        random_codes = manifest.loc[(manifest.stratum == "ngau_nhien") & (manifest.priority_set == stratum), "code"].tolist()
        r = summarize_endpoints(random_codes, final)
        p = summaries[stratum]
        for i, endpoint in enumerate(("E1", "E2", "E3")):
            estimate = r["counts"][i] + pool * p["counts"][i] / p["n"] if pool is not None else math.nan
            estimate_text = "chưa có dữ liệu" if p['missing'][i] == p['n'] else f"{estimate:.2f}"
            lines.append(f"| {stratum} | {endpoint} | {r['counts'][i]} | {pool} × {p['counts'][i]}/{p['n']} | {estimate_text} | {r['missing'][i]} / {p['missing'][i]} |")
    lines += ["", "## Đồng thuận vòng 1", "",
              "Bootstrap theo câu (percentile 95%, " + str(bootstraps) + f" lần, seed {seed}). "
              "Chỉ dùng cặp có cả hai giá trị hợp lệ; báo thiếu theo từng trường. "
              "Cờ lỗi: 'Có' là dương; ô trống là 'Không' khi người duyệt đã chọn tình trạng tổng quát, "
              "và là thiếu khi chưa chọn (HUONG_DAN.md mục 3). "
              "κ không xác định nếu chỉ có một loại ở cả hai người; CI κ bỏ lượt bootstrap không xác định.", "",
              "Mẫu duyệt đôi được làm giàu nghi vấn, không đại diện benchmark. Nhóm ngẫu nhiên dự kiến 12 cặp: **n rất nhỏ**, CI không ổn định.", "",
              "| Nhóm | Trường | Cặp đủ / thiếu | Đồng ý thô (CI) | κ (CI) | Bootstrap κ không xác định |",
              "|---|---|---|---|---|---:|"]
    agreements = {}
    double = manifest[manifest.duyet_doi == "1"]
    for name, group in (("tất cả", double), ("ngẫu nhiên", double[double.stratum == "ngau_nhien"])):
        for field in ["answer", R1_STATUS] + [f[0] for f in DEFECT_FLAGS]:
            pairs = []
            for code in group.code:
                pair = []
                for reviewer in (1, 2):
                    row = records.get((1, reviewer, code), {})
                    if field == "answer":
                        value = row.get("answer")
                    elif field == R1_STATUS:
                        value = clean(row.get(field)) or None
                    else:
                        value = flag_value(row, field, R1_STATUS)
                    pair.append(value if row.get("valid") else None)
                pairs.append(tuple(pair))
            a = agreement(pairs, seed, bootstraps)
            agreements[name, field] = a
            k = f"{a['kappa']:.3f}" if math.isfinite(a['kappa']) else "không xác định"
            lines.append(f"| {name} | {'tập đáp án' if field == 'answer' else field} | {a['n']} / {a['missing']} | {pct(a['raw'])} {ci_text(a['raw_ci'])} | {k} {ci_text(a['kappa_ci'], False)} | {a['undefined_bootstraps']} |")
    lines += ["", "## Thay đổi sau khi mở đáp án", "",
              "So tập đáp án trong cùng người duyệt; mẫu số là câu có đáp án hợp lệ ở cả hai vòng. "
              "Mốc khoá/đa số được xét theo việc tập đáp án có chứa chữ của mốc. Khi hoà đa số, "
              "dùng tất cả chữ đồng hạng và báo riêng. Không tách được ảnh hưởng của khoá, model, giải thích hay việc nghĩ lại; không quy nguyên nhân.", ""]
    changes = {}
    for reviewer in (1, 2):
        codes = manifest.loc[(manifest.duyet_doi == "1") | (reviewer == 1), "code"]
        usable, changed, tied = 0, 0, 0
        directions = {"khoá": Counter(), "đa số 7 model": Counter()}
        for code in codes:
            r1, r2 = records.get((1, reviewer, code), {}), records.get((2, reviewer, code), {})
            if not all(r.get("valid") and r.get("answer") is not None for r in (r1, r2)):
                continue
            usable += 1
            before, after = r1["answer"], r2["answer"]
            if before == after:
                continue
            changed += 1
            key = clean(r2.get("Đáp án chuẩn"))
            key_ref = (key,) if re.fullmatch("[A-Z]", key) else None
            majority = majority_reference(r2)
            tied += bool(majority and len(majority) > 1)
            directions["khoá"][change_direction(before, after, key_ref)] += 1
            directions["đa số 7 model"][change_direction(before, after, majority)] += 1
        changes[reviewer] = dict(n=usable, changed=changed, missing=len(codes)-usable, directions=directions)
        lines += [f"Người {reviewer}: {changed}/{usable} câu đổi ({pct(changed/usable) if usable else 'chưa có dữ liệu'}); "
                  f"{len(codes)-usable} câu thiếu cặp vòng; {tied} câu đổi có đa số hoà.", ""]
        for reference, counts in directions.items():
            lines.append(f"- So với {reference}: " + "; ".join(f"{k}: {counts[k]}" for k in ("về phía mốc", "rời mốc", "giữ quan hệ với mốc", "thiếu mốc")))
        lines.append("")
    lines += ["## Kiểm tra đầu vào", "", f"{len(problems)} vấn đề. Dòng lỗi không dùng làm nhãn; mã vẫn giữ trong mẫu số.", "",
              "| Mã câu | File | Vấn đề |", "|---|---|---|"]
    lines += [f"| {safe(code)} | {safe(file)} | {safe(message)} |" for code, file, message in problems]
    lines += ["", "## Nhãn cuối theo mã", "", "| Mã câu | Khoá | Tập đáp án | Tình trạng | Cờ |", "|---|---|---|---|---|"]
    for code, label in final.items():
        lines.append(f"| {code} | {label['key'] or 'thiếu'} | {show_answer(label['answer'])} | {label['status'] or 'thiếu'} | {label['flag']} |")
    lines += ["", "Không báo recall. phan_xu.xlsx chỉ được đọc, không bao giờ ghi đè. "
              "Mẫu phan_xu_template.xlsx liệt kê mọi bất đồng kể cả đã phân xử; sao chép thành phan_xu.xlsx, "
              "điền bốn trường quyết định, Nguồn, Người phân xử và Lý do. Ô quyết định trống vẫn là chưa phân xử.", ""]
    report = directory / "REVIEW_RESULTS.md"
    report.write_text("\n".join(lines), encoding="utf-8")
    return dict(manifest=manifest, records=records, final=final, disagreements=disagreements,
                problems=problems, summaries=summaries, agreements=agreements, changes=changes, report=report)


def simulate(source, destination, seed=42):
    """Copy only input files, then fill copies. Never load/save source workbooks."""
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("simulation destination must differ from source")
    destination.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source / "sample_manifest.csv", destination / "sample_manifest.csv")
    rng = np.random.default_rng(seed)
    for filename, sheet, round_, reviewer in FILES:
        path = destination / filename
        shutil.copy2(source / filename, path)
        wb = load_workbook(path)
        ws = wb[sheet]
        columns = {c.value: c.column for c in ws[1]}
        fields = ROUND1_FIELDS if round_ == 1 else ROUND2_FIELDS
        for cells in list(ws.iter_rows(min_row=2)):
            number = cells[0].row
            code = str(ws.cell(number, columns[CODE]).value)
            index = int(re.search(r"\d+", code).group())
            options = clean(ws.cell(number, columns[OPTIONS]).value).splitlines()
            n = len(options)
            for field, _, _ in fields:
                ws.cell(number, columns[field]).value = None
            if index % 17 == 0:
                continue
            letter = chr(65 + index % n)
            if reviewer == 2 and index % 3 == 0:
                letter = chr(65 + (index + 1) % n)
            choice, letters = letter, None
            if index % 7 == 0 and n > 1:
                choice, letters = MULTI, "ba"
            elif index % 11 == 0:
                choice = UNKNOWN
            elif index % 19 == 0:
                choice = NONE
            if round_ == 2 and index % 5 == 0:
                choice, letters = chr(65 + (index + reviewer) % n), None
            answer_field, set_field = (R1_ANSWER, R1_SET) if round_ == 1 else (R2_ANSWER, R2_SET)
            ws.cell(number, columns[answer_field], choice)
            ws.cell(number, columns[set_field]).value = letters
            status = STATUS_CHOICES[2 if index % 19 == 0 else 1 if index % 7 == 0 else 0]
            ws.cell(number, columns[R1_STATUS if round_ == 1 else R2_STATUS], status)
            for field, _ in DEFECT_FLAGS:
                if rng.random() < .15:
                    ws.cell(number, columns[field], "Có")
            if round_ == 2:
                key = clean(ws.cell(number, columns["Đáp án chuẩn"]).value)
                key_status = KEY_UNKNOWN if choice == UNKNOWN else KEY_MULTI if choice == MULTI else KEY_OK if choice == key else KEY_WRONG
                ws.cell(number, columns[KEY_STATUS], key_status)
            # Deliberately invalid dropdown, orphan set, absent multi set, and missing status.
            if index % 29 == 0:
                ws.cell(number, columns[answer_field], "giá trị thử không hợp lệ")
            if index % 31 == 0:
                ws.cell(number, columns[answer_field], MULTI)
                ws.cell(number, columns[set_field]).value = None
            if index % 37 == 0:
                ws.cell(number, columns[answer_field], "A")
                ws.cell(number, columns[set_field], "AC")
            if index % 23 == 0:
                ws.cell(number, columns[R1_STATUS if round_ == 1 else R2_STATUS]).value = None
        wb.save(path)
        wb.close()
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dir", type=Path, default=OUT_DIR)
    parser.add_argument("--simulate", action="store_true")
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    if args.simulate:
        # mkdtemp retains the synthetic report for inspection after the process exits.
        directory = Path(tempfile.mkdtemp(prefix="vm14k_review_sim_"))
        simulate(args.dir, directory, args.seed)
    else:
        directory = args.dir
    try:
        result = analyze(directory, args.seed)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"Không thể phân tích: {exc}\n")
    state = "có dữ liệu (xem số đầy đủ/thiếu trong báo cáo)" if any(r.get("filled") for r in result["records"].values()) else "chưa có dữ liệu duyệt"
    print(f"{state}; {len(result['problems'])} vấn đề; {len(result['disagreements'])} bất đồng")
    print(result["report"])


if __name__ == "__main__":
    main()
