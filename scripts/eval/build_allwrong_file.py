#!/usr/bin/env python3
"""
build_allwrong_file.py — the review file for the questions every model gets wrong
(meeting 2026-10-02, action item 1), plus a first analysis of why the models miss them.

Read-only with respect to the dataset and runs. Writes to reports/eval/review/:

    VM14K_125_cau_7_model_deu_sai.xlsx   file for the professors
        HuongDan        how the list was built, how to fill it in
        TomTat_PhanTich findings in words, which models have reasoning, every analysis table
        1_TuTraLoi      optional: question + options only, answer before looking at sheet 2
        2_DoiChieu      key, every model's answer, explanations, automatic hints, reviewer columns
        3_LyDoDayDu     every model's full explanation, one paragraph per row
        4_SuyNghiDayDu  thinking traces (models run with thinking on), for reference
    PHAN_TICH_CAU_TOAN_BO_SAI.md        the summary tables as text, for the team (the workbook is self-contained)

"All wrong" = none of the 7 fully-run models (see build_review_workbook.COUNTED) matches the key
on the frozen test split. Hints are automatic heuristics, not conclusions.

    python scripts/eval/build_allwrong_file.py [--force]
"""
from __future__ import annotations

import argparse
import collections
import difflib
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from scipy.stats import fisher_exact

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_workbook import (  # noqa: E402
    COUNTED, FILL_REVIEW, FONT, FONT_BODY, MODELS, OUT_DIR, REPO_ROOT, WRAP_TOP, clip, fill_letter,
    load_reasoning, load_rows, load_run, option_text, options_block, style_header,
)

XLSX = OUT_DIR / "VM14K_125_cau_7_model_deu_sai.xlsx"
REPORT = OUT_DIR / "PHAN_TICH_CAU_TOAN_BO_SAI.md"
EXPLAINER = "Nemotron-3-Ultra"
STRONG_VOTES = 5   # of 7 models on the same wrong option
SCATTERED_VOTES = 3  # 7 votes over ≤3 wrong options: the top one always has ≥3

# Question features. Deliberately simple regexes: they are hints for reviewers and inputs to an
# exploratory comparison, so a short, readable rule beats a clever one.
FEATURES = [
    ("Hỏi phủ định / loại trừ",
     lambda r: bool(re.search(r"\b(sai|ngoại trừ|không đúng|không phải|không có|không thuộc|không nên|"
                              r"chống chỉ định|TRỪ|KHÔNG)\b|,\s*trừ\b", r["question"]))),
    ("Hỏi đếm số lượng (bao nhiêu / mấy)", lambda r: bool(re.search(r"(?i)\b(bao nhiêu|mấy)\b", r["question"]))),
    ("Phương án là số liệu / ngưỡng",
     lambda r: sum(bool(re.search(r"\d", o)) for o in r["options"]) >= max(2, len(r["options"]) - 1)),
    ("Nhắc tới hình / bảng / ảnh",
     lambda r: bool(re.search(r"(?i)\b(hình|ảnh|bảng|sơ đồ|biểu đồ|đây là|phim|chi tiết số)\b", r["question"]))),
    ("Phương án tham chiếu phương án khác (Cả A và B…)",
     lambda r: any(re.search(r"(?i)\b(cả|tất cả)\b.*\b(đều|trên|đúng|sai)\b|\b[A-Ga-g]\s*(,|và|\+)\s*[A-Ga-g]\b"
                             r"|\bcâu\s+[a-g]\b", o) for o in r["options"])),
    ("Có hai phương án gần giống nhau",
     lambda r: any(difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio() > 0.85
                   for i, a in enumerate(r["options"]) for b in r["options"][i + 1:] if len(a) > 15)),
    ("Câu đúng / sai (2 phương án)", lambda r: len(r["options"]) == 2),
    ("Câu hỏi rất ngắn (< 30 ký tự)", lambda r: len(r["question"].strip()) < 30),
    ("Liên quan quy định / chương trình y tế VN",
     lambda r: bool(re.search(r"(?i)bộ y tế|thông tư|quyết định|luật|nghị định|quốc gia|chương trình|"
                              r"tiêm chủng|bảo hiểm|trạm y tế", r["question"] + " " + " ".join(r["options"])))),
]

HINTS = {  # order = display priority for the professors
    "nghi_khoa_sai": "Nghi đáp án chuẩn sai: ≥5/7 model cùng chọn một đáp án khác, Nemotron giữ đáp án đó khi lập luận",
    "nghi_cau_loi": "Nghi câu hỏi lỗi: Nemotron từ chối chọn (thiếu hình / thiếu dữ kiện / không có đáp án đúng)",
    "dong_thuan_manh": "≥5/7 model cùng chọn một đáp án khác khoá (Nemotron không giữ khi lập luận)",
    "dong_thuan_vua": "4/7 model cùng chọn một đáp án khác khoá",
    "phan_tan": "Model chọn phân tán (đáp án sai phổ biến nhất chỉ 3/7): câu khó hoặc mơ hồ",
    "sua_duoc": "Nemotron chọn đúng khoá khi được yêu cầu lập luận: lỗi có thể do trả lời nhanh",
}

REVIEW_FIELDS = [
    ("Phân loại (Thầy/Cô)", 30, ["Đáp án chuẩn sai", "Câu hỏi lỗi / mơ hồ / thiếu dữ kiện", "Nhiều đáp án đúng",
                                 "Đúng theo phác đồ / thực hành VN", "Kiến thức cũ / đã thay đổi",
                                 "Model sai, đáp án chuẩn đúng", "Không chắc"]),
    ("Đáp án đúng theo Thầy/Cô", 14, ["A", "B", "C", "D", "E", "F", "G", "Nhiều đáp án đúng",
                                       "Không có đáp án đúng", "Không xác định được"]),
    ("Nguồn / bằng chứng", 34, None),
    ("Ghi chú", 34, None),
    ("Người duyệt", 12, None),
]
ROUND1_FIELDS = [
    ("Đáp án Thầy/Cô chọn", 14, REVIEW_FIELDS[1][2]),
    ("Mức chắc chắn", 12, ["Chắc chắn", "Khá chắc", "Không chắc"]),
    ("Ghi chú", 34, None),
]


OUTCOMES = ["Chọn đúng khoá", "Giữ đáp án sai của đa số", "Chọn một đáp án sai khác", "Từ chối chọn"]
VN_SOURCE = re.compile(r"(?i)bộ y tế|\bbyt\b|phác đồ|việt nam|thông tư|quyết định số|giáo trình|đại học y|"
                       r"hướng dẫn chẩn đoán và điều trị")
INTL_SOURCE = re.compile(r"(?i)\bwho\b|\baha\b|\besc\b|\bgina\b|\bgold\b|kdigo|nccn|acog|\bnice\b|\bcdc\b|harrison|"
                         r"uptodate|williams|robbins|guyton|nelson|medscape|\bfda\b|idsa|\bada\b|easl|aasld")


def outcome_of(rec, key, top):
    """How an explanation ends, relative to the key and to the models' majority wrong answer."""
    if rec.get("status") == "abstained" or not rec.get("pred"):
        return OUTCOMES[3]
    if rec["pred"] == key:
        return OUTCOMES[0]
    return OUTCOMES[1] if rec["pred"] in top else OUTCOMES[2]


def cited_sources(text):
    """Which kind of reference an explanation invokes — a rough hint for the Vietnam-practice question."""
    vn, intl = bool(VN_SOURCE.search(text or "")), bool(INTL_SOURCE.search(text or ""))
    return "Cả VN và quốc tế" if vn and intl else "Nhắc nguồn / bối cảnh VN" if vn else \
        "Chỉ nguồn quốc tế" if intl else "Không nêu nguồn cụ thể"


def conclusion(text, cap=900):
    """The explanation from its last "Kết luận" line on (307/320 have one), else its last two paragraphs.
    Full explanations are ~2,100 characters, too tall for one Excel row, so sheet 2 shows this part."""
    text = (text or "").strip()
    lines = text.splitlines()
    idx = max((i for i, line in enumerate(lines) if "kết luận" in line.lower()), default=None)
    if idx is None:
        part = "\n\n".join([p for p in re.split(r"\n\s*\n", text) if p.strip()][-2:])
    else:
        part = "\n".join(lines[idx:]).strip()
    if len(part) > cap:
        part = part[:cap].rsplit(" ", 1)[0] + " … (xem toàn văn ở sheet 3_LyDoDayDu)"
    return part


def add_fields(ws, fields, first_col, n_rows):
    for k, (_, _, choices) in enumerate(fields):
        col = first_col + k
        if choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{n_rows + 1}")
        for r in range(2, n_rows + 2):
            ws.cell(row=r, column=col).fill = FILL_REVIEW
            ws.cell(row=r, column=col).alignment = WRAP_TOP


def write_guide(ws, n_wrong, n_test):
    ws.column_dimensions["A"].width = 125
    lines = [
        ("VM14K — các câu mà cả 7 model đều trả lời sai", True),
        ("", False),
        (f"Có {n_wrong}/{n_test} câu trong tập test (sau làm sạch) mà cả 7 model đều chọn khác đáp án chuẩn: "
         + ", ".join(COUNTED) + ". Các model trả lời zero-shot, chỉ một chữ cái, theo đúng prompt của bài báo VM14K.", False),
        ("Câu hỏi chúng em muốn nhờ Thầy/Cô trả lời: với mỗi câu, ai đúng — đáp án chuẩn của bộ đề hay các model? "
         "Nguyên nhân có thể là: đáp án chuẩn sai; câu hỏi lỗi, mơ hồ hoặc thiếu dữ kiện; có nhiều đáp án đúng; "
         "đáp án đúng theo phác đồ hoặc thực hành tại Việt Nam mà model không biết; kiến thức cũ; hoặc model thật sự sai.", False),
        ("", False),
        ("Cách làm (khuyến nghị):", True),
        ("1. Nếu Thầy/Cô có thời gian: làm sheet '1_TuTraLoi' TRƯỚC. Sheet này chỉ có câu hỏi và các phương án; "
         "Thầy/Cô tự chọn đáp án như khi làm đề. Làm vậy để ý kiến của Thầy/Cô không bị đáp án chuẩn hay đáp án model dẫn dắt.", False),
        ("2. Sheet '2_DoiChieu': mỗi câu có đáp án chuẩn, đáp án của từng model (xanh = trùng đáp án chuẩn, đỏ = khác), "
         "đáp án sai được nhiều model chọn nhất, đoạn kết luận trong lời giải thích tiếng Việt của Nemotron-3-Ultra, và "
         "đáp án các model khác chốt khi được yêu cầu giải thích (gpt-oss-20b và các model chạy trên máy). "
         "Sheet '3_LyDoDayDu' có toàn văn lời giải thích, mỗi đoạn một dòng. Thầy/Cô điền 5 cột màu vàng ở cuối mỗi dòng (cột 'Phân loại' và 'Đáp án đúng' có danh sách chọn sẵn).", False),
        ("3. Các câu được xếp theo 'Nhóm gợi ý' — nhóm nghi đáp án chuẩn sai lên đầu. Nhóm gợi ý và cột 'Đặc điểm câu' do máy tự gán "
         "theo quy tắc đơn giản, CHỈ để tham khảo, không phải kết luận.", False),
        ("4. Sheet 'TomTat_PhanTich' (ngay sau sheet này): tóm tắt kết quả, danh sách các model đã có lời giải thích, "
         "và toàn bộ bảng phân tích vì sao model sai. Mọi thứ nằm trong file này, không cần mở file nào khác.", False),
        ("", False),
        ("Lưu ý:", True),
        ("• Model cùng sai chưa có nghĩa đáp án chuẩn sai — nhiều model có thể cùng học một kiến thức sai.", False),
        ("• Lời giải thích của model là lời model tự nói, có thể sai hoặc bịa nguồn; chỉ dùng để gợi ý.", False),
        ("• Nếu một câu có nhiều đáp án đúng, xin chọn 'Nhiều đáp án đúng' và ghi các chữ vào cột 'Ghi chú'.", False),
        ("• Mã câu (C001…) dùng chung cho các sheet; cột 'id' ở cuối sheet 2 là mã nội bộ của bộ dữ liệu.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = Font(name=FONT, bold=bold, size=14 if ws.max_row == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--force", action="store_true", help="overwrite the workbook (it may already be filled in)")
    args = ap.parse_args()
    if XLSX.exists() and not args.force:
        sys.exit(f"{XLSX.relative_to(REPO_ROOT)} exists and may hold reviewer input — pass --force to overwrite")

    rows = load_rows()
    runs = {label: load_run(stem) for label, stem, _ in MODELS}
    reasoning, _ = load_reasoning()
    explain = reasoning.get(EXPLAINER, {})
    test_ids = sorted(runs[COUNTED[0]])
    wrong = [i for i in test_ids if not any(runs[m][i]["correct"] for m in COUNTED)]
    # Every other model with explanations for these questions, in a fixed order. Hints stay based on
    # Nemotron (the only explainer covering all questions from the start); the others are extra evidence.
    order = ["gpt-oss-20b"] + [label for label, _, _ in MODELS]
    others = {label: reasoning[label] for label in sorted(reasoning, key=lambda l: order.index(l) if l in order else 99)
              if label != EXPLAINER and any(q in reasoning[label] for q in wrong)}

    info = {}
    for qid in wrong:
        row = rows[qid]
        votes = collections.Counter(runs[m][qid]["pred"] for m in COUNTED if runs[m][qid]["pred"])
        top_n = max(votes.values()) if votes else 0
        top = sorted(l for l, n in votes.items() if n == top_n)
        rec = explain.get(qid)
        exp_pred = rec["pred"] if rec else None
        if rec is None:
            exp_outcome = "Chưa có giải thích"
        elif rec.get("status") == "abstained":
            exp_outcome = "Từ chối chọn"
        elif exp_pred == row["answer"]:
            exp_outcome = "Chọn đúng khoá khi lập luận"
        elif exp_pred in top:
            exp_outcome = "Giữ đáp án sai của đa số"
        else:
            exp_outcome = "Chọn một đáp án sai khác"
        if exp_outcome == "Từ chối chọn":
            hint = "nghi_cau_loi"
        elif exp_outcome == "Chọn đúng khoá khi lập luận":
            hint = "sua_duoc"
        elif top_n >= STRONG_VOTES:
            hint = "nghi_khoa_sai" if exp_outcome == "Giữ đáp án sai của đa số" else "dong_thuan_manh"
        elif top_n <= SCATTERED_VOTES:
            hint = "phan_tan"
        else:
            hint = "dong_thuan_vua"
        explained = {EXPLAINER: rec} if rec else {}
        explained.update({label: recs[qid] for label, recs in others.items() if qid in recs})
        tally = collections.Counter(outcome_of(r, row["answer"], top) for r in explained.values())
        info[qid] = dict(votes=votes, top=top, top_n=top_n, rec=rec, exp_outcome=exp_outcome, hint=hint,
                         features=[name for name, f in FEATURES if f(row)], explained=explained, tally=tally,
                         sources=cited_sources(rec.get("answer_text") if rec else ""))

    hint_order = list(HINTS)
    wrong.sort(key=lambda q: (hint_order.index(info[q]["hint"]), -info[q]["top_n"]))
    codes = {qid: f"C{n:03d}" for n, qid in enumerate(wrong, start=1)}

    wb = Workbook()
    write_guide(wb.active, len(wrong), len(test_ids))
    wb.active.title = "HuongDan"

    # Summary first, so the professors need no other file: findings in words, which models have
    # reasoning, then every analysis table.
    stats = analysis(rows, runs, test_ids, wrong, info)
    status = model_status(reasoning, wrong)
    ws = wb.create_sheet("TomTat_PhanTich")
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18
    for text, bold in summary_lines(rows, test_ids, wrong, info, status):
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = Font(name=FONT, bold=bold, size=14 if ws.max_row == 1 else 11)
        cell.alignment = Alignment(wrap_text=False, vertical="top")
    ws.append([])
    status_block = ("Các model đã có lời giải thích (reasoning) cho các câu này",
                    ["Model", "Chạy ở đâu", "Có phần suy nghĩ (thinking)?", "Số câu có giải thích", "Trạng thái"],
                    status, f"Cập nhật lúc {datetime.now():%H:%M %d/%m/%Y}. Bảng tự cập nhật mỗi lần tạo lại file.")
    for block_title, header, table, note in [status_block] + stats:
        ws.append([block_title])
        ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True, size=12)
        ws.append(header)
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, bold=True)
            c.fill = FILL_REVIEW
        for line in table:
            ws.append(line)
            for c in ws[ws.max_row]:
                c.font = FONT_BODY
        if note:
            ws.append([note])
            ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, italic=True, size=9)
        ws.append([])

    # 1 — blind: question and options only
    ws = wb.create_sheet("1_TuTraLoi")
    base = [("Mã câu", 8), ("Chuyên khoa", 20), ("Câu hỏi", 60), ("Các phương án", 60)]
    style_header(ws, [h for h, _ in base] + [h for h, _, _ in ROUND1_FIELDS],
                 [w for _, w in base] + [w for _, w, _ in ROUND1_FIELDS])
    for qid in wrong:
        row = rows[qid]
        ws.append([codes[qid], ", ".join(row["medical_topic"]), row["question"], options_block(row)]
                  + [""] * len(ROUND1_FIELDS))
        for cell in ws[ws.max_row]:
            cell.font, cell.alignment = FONT_BODY, WRAP_TOP
    add_fields(ws, ROUND1_FIELDS, len(base) + 1, len(wrong))
    ws.auto_filter.ref = ws.dimensions

    # 2 — reveal: key, models, explanation, hints, reviewer columns
    ws = wb.create_sheet("2_DoiChieu")
    base = ([("Mã câu", 8), ("Nhóm gợi ý (tự động)", 30), ("Chuyên khoa", 18), ("Độ khó (nhãn)", 10),
             ("Câu hỏi", 50), ("Các phương án", 52), ("Đáp án chuẩn", 8), ("Nội dung đáp án chuẩn", 28),
             ("Đáp án sai được chọn nhiều nhất", 10), ("Nội dung đáp án đó", 28), ("Số model chọn (/7)", 8)]
            + [(label, 9) for label, _, _ in MODELS]
            + [("Nemotron khi lập luận", 11), ("Nemotron: lý do — đoạn kết luận (toàn văn ở sheet 3)", 70)]
            + [(f"{label} khi lập luận", 11) for label in others]
            + [("Khi lập luận: đúng khoá / giữ sai đa số / sai khác / từ chối", 16),
               ("Nemotron viện dẫn", 16), ("Đặc điểm câu (tự động)", 30)])
    tail = [("id", 12)]
    style_header(ws, [h for h, _ in base] + [h for h, _, _ in REVIEW_FIELDS] + [h for h, _ in tail],
                 [w for _, w in base] + [w for _, w, _ in REVIEW_FIELDS] + [w for _, w in tail])
    model_col0 = 12
    for qid in wrong:
        row, d = rows[qid], info[qid]
        gold = row["answer"]
        letters = [runs[label].get(qid, {}).get("pred") for label, _, _ in MODELS]
        rec = d["rec"]
        ws.append([codes[qid], HINTS[d["hint"]], ", ".join(row["medical_topic"]), row["difficulty_level"],
                   row["question"], options_block(row), gold, option_text(row, gold),
                   " / ".join(d["top"]) + (" (hoà)" if len(d["top"]) > 1 else ""),
                   "\n".join(f"{l}. {option_text(row, l)}" for l in d["top"]), d["top_n"]]
                  + [l or "" for l in letters]
                  + [(rec["pred"] or "Không chọn") if rec else "", conclusion(rec.get("answer_text")) if rec else ""]
                  + [((recs[qid]["pred"] or "Không chọn") if qid in recs else "") for recs in others.values()]
                  + [" / ".join(str(d["tally"][o]) for o in OUTCOMES) + f" (của {sum(d['tally'].values())} model)",
                     d["sources"], "; ".join(d["features"])]
                  + [""] * len(REVIEW_FIELDS) + [qid])
        r = ws.max_row
        for cell in ws[r]:
            cell.font, cell.alignment = FONT_BODY, WRAP_TOP
        for k, letter in enumerate(letters):
            fill_letter(ws.cell(row=r, column=model_col0 + k), letter, gold)
        if rec:
            fill_letter(ws.cell(row=r, column=model_col0 + len(MODELS)), rec["pred"], gold)
        for k, recs in enumerate(others.values()):
            if qid in recs:
                fill_letter(ws.cell(row=r, column=model_col0 + len(MODELS) + 2 + k), recs[qid]["pred"], gold)
    add_fields(ws, REVIEW_FIELDS, len(base) + 1, len(wrong))
    ws.auto_filter.ref = ws.dimensions

    # 3 — full explanations, one paragraph per row so no row exceeds Excel's height limit
    ws = wb.create_sheet("3_LyDoDayDu")
    style_header(ws, ["Mã câu", "Model", "Lời giải thích (toàn văn, mỗi đoạn một dòng)"], [8, 18, 125])
    for qid in wrong:
        ws.append([codes[qid], "", f"{rows[qid]['question']}  (đáp án chuẩn: {rows[qid]['answer']})"])
        for cell in ws[ws.max_row]:
            cell.font, cell.alignment = Font(name=FONT, bold=True, size=11), WRAP_TOP
        for label, rec in info[qid]["explained"].items():
            ws.append([codes[qid], label, f"— {label} chọn: {rec['pred'] or 'không chọn'} —"])
            for cell in ws[ws.max_row]:
                cell.font, cell.alignment = Font(name=FONT, bold=True, size=10), WRAP_TOP
            fill_letter(ws.cell(row=ws.max_row, column=3), rec["pred"], rows[qid]["answer"])
            ws.cell(row=ws.max_row, column=3).alignment = WRAP_TOP
            for para in [p.strip() for p in re.split(r"\n\s*\n", rec.get("answer_text") or "") if p.strip()]:
                ws.append([codes[qid], label, para])
                for cell in ws[ws.max_row]:
                    cell.font, cell.alignment = FONT_BODY, WRAP_TOP
    ws.auto_filter.ref = ws.dimensions

    # 4 — full thinking traces
    ws = wb.create_sheet("4_SuyNghiDayDu")
    style_header(ws, ["Mã câu", "Model", "Chọn", "Phần suy nghĩ (thinking, thường bằng tiếng Anh)"], [8, 18, 10, 130])
    for qid in wrong:
        for label, rec in info[qid]["explained"].items():
            if not rec.get("reasoning"):
                continue  # local models ran with thinking off: no trace
            ws.append([codes[qid], label, rec["pred"] or "Không chọn", clip(rec["reasoning"])])
            for cell in ws[ws.max_row]:
                cell.font, cell.alignment = FONT_BODY, WRAP_TOP

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    (OUT_DIR / "allwrong_ids.txt").write_text("\n".join(wrong) + "\n", encoding="utf-8")
    write_report([status_block] + stats, len(wrong), len(test_ids))
    print(f"{len(wrong)} questions; hints: "
          + ", ".join(f"{k} {sum(info[q]['hint'] == k for q in wrong)}" for k in HINTS))
    print(f"wrote {XLSX.relative_to(REPO_ROOT)} and {REPORT.relative_to(REPO_ROOT)}")


def analysis(rows, runs, test_ids, wrong, info):
    """[(title, header, rows, note)] — the 'why do all models miss' tables."""
    n_w, n_t = len(wrong), len(test_ids)
    wrong_set = set(wrong)
    pct = lambda k, n: f"{100 * k / n:.1f}%"
    blocks = []

    count = collections.Counter(info[q]["top_n"] for q in wrong)
    blocks.append(("Mức hội tụ: số model (trong 7) cùng chọn một đáp án sai", ["Số model", "Số câu", "Tỉ lệ"],
                   [[k, count[k], pct(count[k], n_w)] for k in sorted(count, reverse=True)],
                   "Hội tụ cao = nhiều model cùng một đáp án khác khoá → đáng nghi khoá sai; phân tán = câu khó/mơ hồ. Với 4 phương án (3 đáp án sai), mức thấp nhất có thể là 3/7."))

    count = collections.Counter(info[q]["exp_outcome"] for q in wrong)
    blocks.append(("Nemotron-3-Ultra khi được yêu cầu lập luận (thinking bật, giải thích tiếng Việt)",
                   ["Kết quả", "Số câu", "Tỉ lệ"],
                   [[k, v, pct(v, n_w)] for k, v in count.most_common()],
                   "Lần trả lời nhanh của Nemotron không được tính vào 7 model (chạy chưa đủ tập test)."))

    count = collections.Counter(info[q]["hint"] for q in wrong)
    blocks.append(("Nhóm gợi ý (tự động)", ["Nhóm", "Số câu", "Tỉ lệ"],
                   [[HINTS[k], count[k], pct(count[k], n_w)] for k in HINTS], None))

    per_model = collections.defaultdict(collections.Counter)
    for q in wrong:
        for label, rec in info[q]["explained"].items():
            per_model[label][outcome_of(rec, rows[q]["answer"], info[q]["top"])] += 1
    blocks.append(("Từng model khi được yêu cầu giải thích rồi chốt đáp án",
                   ["Model", "Số câu có giải thích"] + OUTCOMES,
                   [[label, sum(c.values())] + [f"{c[o]} ({pct(c[o], sum(c.values()))})" for o in OUTCOMES]
                    for label, c in per_model.items()],
                   "Nemotron và gpt-oss-20b bật thinking (qua NVIDIA). Các model local chạy trên CPU với thinking tắt "
                   "nên chỉ có phần giải thích; model nào chưa chạy xong thì số câu ít hơn 125."))

    # The two explainers that ran with thinking on and cover every question: where do they agree?
    second = "gpt-oss-20b"
    if all(second in info[q]["explained"] and EXPLAINER in info[q]["explained"] for q in wrong):
        cross = collections.Counter(
            (outcome_of(info[q]["explained"][EXPLAINER], rows[q]["answer"], info[q]["top"]),
             outcome_of(info[q]["explained"][second], rows[q]["answer"], info[q]["top"])) for q in wrong)
        blocks.append((f"{EXPLAINER} (hàng) × {second} (cột) khi lập luận — số câu",
                       [f"{EXPLAINER} \\ {second}"] + OUTCOMES,
                       [[o1] + [cross[o1, o2] for o2 in OUTCOMES] for o1 in OUTCOMES],
                       "Đường chéo = hai model cùng kết luận. Cả hai cùng giữ đáp án sai của đa số → nghi khoá sai mạnh hơn; "
                       "cả hai cùng chọn đúng khoá → nhiều khả năng lỗi do trả lời nhanh; cả hai cùng từ chối → nghi câu hỏi lỗi."))

    count = collections.Counter(info[q]["sources"] for q in wrong)
    kept = collections.Counter(info[q]["sources"] for q in wrong if info[q]["exp_outcome"] == "Giữ đáp án sai của đa số")
    blocks.append(("Nguồn Nemotron viện dẫn trong lời giải thích",
                   ["Loại nguồn", "Số câu", "…trong đó Nemotron giữ đáp án sai của đa số"],
                   [[k, f"{v} ({pct(v, n_w)})", kept[k]] for k, v in count.most_common()],
                   "Dò từ khoá đơn giản (Bộ Y tế, phác đồ, Việt Nam… / WHO, AHA, Harrison…). Câu model giữ đáp án sai "
                   "trong khi viện dẫn nguồn quốc tế là ứng viên cho nhóm 'đúng theo thực hành VN', cần chuyên gia xác nhận."))

    blocks.append(("Các giả thuyết đã kiểm tra (không cần chuyên gia)", ["Giả thuyết", "Kết quả", "Kết luận"],
                   hypotheses(rows, runs, test_ids, wrong, info), None))

    feature_rows = []
    for name, f in FEATURES:
        a = sum(f(rows[q]) for q in wrong)
        b = sum(f(rows[q]) for q in test_ids)
        rest_yes, rest_no = b - a, (n_t - n_w) - (b - a)
        _, p = fisher_exact([[a, n_w - a], [rest_yes, rest_no]])
        lift = (a / n_w) / (b / n_t) if b else float("nan")
        feature_rows.append([name, f"{a} ({pct(a, n_w)})", f"{b} ({pct(b, n_t)})", f"{lift:.2f}", f"{p:.3f}"])
    feature_rows.sort(key=lambda r: -float(r[3]))
    blocks.append(("Đặc điểm câu: nhóm 7 model đều sai so với toàn bộ tập test",
                   ["Đặc điểm", f"Trong {n_w} câu sai", f"Trong {n_t} câu test", "Hệ số (lift)", "p (Fisher)"],
                   feature_rows,
                   "Lift > 1: đặc điểm xuất hiện nhiều hơn bình thường trong nhóm toàn bộ sai. Phân tích khám phá: "
                   "nhiều phép so sánh, quy tắc gán nhãn đơn giản, chưa hiệu chỉnh p — chỉ dùng để gợi hướng."))

    for title, key in [("Theo nhãn độ khó", lambda r: [r["difficulty_level"]]),
                       ("Theo số phương án", lambda r: [str(len(r["options"]))])]:
        cw = collections.Counter(v for q in wrong for v in key(rows[q]))
        ct = collections.Counter(v for q in test_ids for v in key(rows[q]))
        blocks.append((title, ["Nhóm", f"Trong {n_w} câu sai", f"Trong {n_t} câu test", "Tỉ lệ sai toàn bộ"],
                       [[k, f"{cw[k]} ({pct(cw[k], n_w)})", f"{ct[k]} ({pct(ct[k], n_t)})", pct(cw[k], ct[k])]
                        for k in sorted(ct)], None))

    cw = collections.Counter(t for q in wrong for t in rows[q]["medical_topic"])
    ct = collections.Counter(t for q in test_ids for t in rows[q]["medical_topic"])
    topic_rows = [[t, cw[t], ct[t], pct(cw[t], ct[t])] for t in ct if ct[t] >= 15]
    topic_rows.sort(key=lambda r: -r[1] / r[2])
    blocks.append(("Theo chuyên khoa (chuyên khoa có ≥15 câu test)",
                   ["Chuyên khoa", "Số câu cả 7 model sai", "Số câu test", "Tỉ lệ"], topic_rows,
                   f"Tỉ lệ chung: {pct(n_w, n_t)}. Một câu có thể thuộc nhiều chuyên khoa. Cỡ mẫu từng chuyên khoa nhỏ."))
    return blocks


# (label, where it runs, thinking?) for every model whose explanations belong in the file, in display order.
EXPLAINERS = [
    ("Nemotron-3-Ultra", "NVIDIA (API)", "Có"),
    ("gpt-oss-20b", "NVIDIA (API)", "Có"),
    ("DeepSeek-v4.1-flash", "NVIDIA (API)", "Có"),
    ("Gemma-4-31B", "NVIDIA (API)", "Có"),
    ("Qwen3.5-9B", "Máy local (Ollama, CPU)", "Không — chỉ có phần giải thích"),
    ("MedGemma-4B", "Máy local (Ollama, CPU)", "Không — chỉ có phần giải thích"),
    ("Llama-3.1-8B", "Máy local (Ollama, CPU)", "Không — chỉ có phần giải thích"),
    ("Qwen3-8B", "Máy local (Ollama, CPU)", "Không — chỉ có phần giải thích"),
    ("Gemma-4-12B", "Máy local (Ollama, CPU)", "Không — chỉ có phần giải thích"),
]


def model_status(reasoning, wrong):
    """One row per model: how many of the questions have an explanation, and what is still pending."""
    rows_out = []
    for label, where, thinking in EXPLAINERS:
        n = sum(q in reasoning.get(label, {}) for q in wrong)
        if n == len(wrong):
            state = "Đủ"
        elif n:
            state = "Đang chạy"
        elif where.startswith("NVIDIA"):
            state = "Chưa lấy được — NVIDIA đang quá tải, sẽ thử lại"
        else:
            state = "Đang chờ (các model local chạy lần lượt)"
        rows_out.append([label, where, thinking, f"{n}/{len(wrong)}", state])
    return rows_out


def summary_lines(rows, test_ids, wrong, info, status):
    """The findings in words for the top of the summary sheet. Every number is computed here."""
    n_w, n_t = len(wrong), len(test_ids)
    hint = collections.Counter(info[q]["hint"] for q in wrong)
    outcome = collections.Counter(info[q]["exp_outcome"] for q in wrong)
    ready = [r[0] for r in status if r[4] == "Đủ"]
    running = [f"{r[0]} ({r[3]})" for r in status if r[4] == "Đang chạy"]
    level = {lv: (sum(rows[q]["difficulty_level"] == lv for q in wrong),
                  sum(rows[q]["difficulty_level"] == lv for q in test_ids)) for lv in ("Easy", "Medium", "Challenging")}
    two = sum(len(rows[q]["options"]) == 2 for q in wrong), sum(len(rows[q]["options"]) == 2 for q in test_ids)
    both = []
    if all({EXPLAINER, "gpt-oss-20b"} <= info[q]["explained"].keys() for q in wrong):
        agree = collections.Counter(
            outcome_of(info[q]["explained"][EXPLAINER], rows[q]["answer"], info[q]["top"])
            for q in wrong
            if outcome_of(info[q]["explained"][EXPLAINER], rows[q]["answer"], info[q]["top"])
            == outcome_of(info[q]["explained"]["gpt-oss-20b"], rows[q]["answer"], info[q]["top"]))
        both = [(f"• Hai model có thinking ({EXPLAINER} và gpt-oss-20b) cùng kết luận: cùng giữ đáp án sai của đa số ở "
                 f"{agree[OUTCOMES[1]]} câu (nghi khoá sai mạnh nhất), cùng chọn đúng khoá ở {agree[OUTCOMES[0]]} câu, "
                 f"cùng từ chối chọn ở {agree[OUTCOMES[3]]} câu (bảng chéo bên dưới).", False)]
    return [
        ("Tóm tắt kết quả — các câu mà cả 7 model đều trả lời sai", True),
        ("", False),
        (f"• {n_w}/{n_t} câu test ({100 * n_w / n_t:.1f}%) bị cả 7 model chọn khác đáp án chuẩn.", False),
        (f"• Model có lời giải thích đầy đủ cho cả {n_w} câu: {', '.join(ready) or 'chưa có'}"
         + (f"; đang chạy: {', '.join(running)}" if running else "") + " (chi tiết ở bảng bên dưới).", False),
        ("", False),
        ("Phân nhóm tự động (chỉ để ưu tiên duyệt, không phải kết luận):", True),
        (f"• {hint['nghi_khoa_sai']} câu nghi đáp án chuẩn sai: từ 5/7 model trở lên cùng chọn một đáp án khác, và "
         "Nemotron vẫn giữ đáp án đó khi được yêu cầu lập luận.", False),
        (f"• {hint['nghi_cau_loi']} câu nghi câu hỏi lỗi: Nemotron từ chối chọn (thiếu hình / thiếu dữ kiện / không có đáp án đúng).", False),
        (f"• {hint['sua_duoc']} câu ({100 * outcome['Chọn đúng khoá khi lập luận'] / n_w:.0f}%) Nemotron chọn ĐÚNG đáp án chuẩn "
         "khi được phép lập luận: một phần các câu 'toàn bộ sai' có thể do cách chấm chỉ cho trả lời 1 chữ cái.", False),
        (f"• {hint['dong_thuan_manh'] + hint['dong_thuan_vua']} câu các model đồng thuận vào một đáp án khác khoá ở mức "
         f"vừa/mạnh; {hint['phan_tan']} câu model chọn phân tán (câu khó hoặc mơ hồ).", False),
        *both,
        ("", False),
        ("Đặc điểm của các câu này so với toàn bộ tập test (xem bảng 'Đặc điểm câu'):", True),
        ("• Gặp nhiều hơn bình thường: câu hỏi đếm số lượng ('có bao nhiêu…'), có hai phương án gần giống nhau, "
         "phương án là số liệu / ngưỡng.", False),
        (f"• Câu đúng/sai (2 phương án) ít hơn hẳn: {two[0]}/{n_w} so với {two[1]}/{n_t} — vì chỉ có 2 lựa chọn nên khó để 7 model cùng sai.", False),
        ("• Nhãn độ khó không liên quan: tỉ lệ cả 7 model sai là "
         + ", ".join(f"{lv} {100 * a / b:.1f}%" for lv, (a, b) in level.items()) + ".", False),
        ("", False),
        ("Các giả thuyết đã kiểm tra và loại bỏ (xem bảng 'Các giả thuyết đã kiểm tra'):", True),
        ("• Đáp án chuẩn bị lệch một vị trí khi trích đề: không có dấu hiệu.", False),
        ("• Cùng một câu xuất hiện ở chỗ khác trong bộ đề với đáp án khác: không có (0 câu).", False),
        ("• Model chỉ biết nguồn quốc tế: Nemotron hầu như luôn nhắc giáo trình / bối cảnh Việt Nam; chỉ "
         f"{sum(info[q]['sources'] == 'Chỉ nguồn quốc tế' for q in wrong)} câu dẫn riêng nguồn quốc tế.", False),
        ("", False),
        ("Lưu ý: mọi phân nhóm và đặc điểm do máy gán bằng quy tắc đơn giản; lời giải thích là lời model tự nói, có thể sai. "
         "Kết luận cuối cùng cần ý kiến của Thầy/Cô.", False),
    ]


def _words(text):
    """Word/number tokens with diacritics kept (they carry meaning in Vietnamese)."""
    return tuple(re.findall(r"\w+", unicodedata.normalize("NFC", text or "").lower()))


def _option_key(text):
    """Order-insensitive for listed items: 'a, b, c' == 'a, c, b'."""
    parts = [_words(p) for p in re.split(r"[,;]| và ", unicodedata.normalize("NFC", text or "").lower())]
    return tuple(sorted(p for p in parts if p))


def hypotheses(rows, runs, test_ids, wrong, info):
    """Explanations for 'all models wrong' that can be checked without a clinician."""
    letters = "ABCDEFG"
    out = []
    # 1. Key letter shifted by one position (an extraction/labelling slip)?
    four = [q for q in wrong if len(rows[q]["options"]) == 4 and len(info[q]["top"]) == 1]
    adj = sum(abs(letters.index(info[q]["top"][0]) - rows[q]["answer_index"]) == 1 for q in four)
    base_adj = base_n = 0
    wrong_set = set(wrong)
    for m in COUNTED:
        for q in test_ids:
            r = runs[m][q]
            if q in wrong_set or len(rows[q]["options"]) != 4 or not r["pred"] or r["correct"]:
                continue
            base_n += 1
            base_adj += abs(letters.index(r["pred"]) - rows[q]["answer_index"]) == 1
    out.append(["Đáp án chuẩn bị lệch một vị trí (ví dụ đúng là B nhưng khoá ghi A)",
                f"Đáp án đa số nằm cạnh khoá: {adj}/{len(four)} = {100*adj/len(four):.1f}%; mọi câu trả lời sai "
                f"khác: {100*base_adj/base_n:.1f}%; nếu ngẫu nhiên: 48.5%",
                "Không có dấu hiệu lệch vị trí"])
    # 2. Models drawn to the longest option?
    def longest(ids, pick):
        k = n = 0
        for q in ids:
            lens = [len(o) for o in rows[q]["options"]]
            if lens.count(max(lens)) == 1 and pick(q):
                n += 1
                k += lens.index(max(lens)) == letters.index(pick(q))
        return k, n
    k1, n1 = longest(four, lambda q: info[q]["top"][0])
    k2, n2 = longest(four, lambda q: rows[q]["answer"])
    out.append(["Model bị hút về phương án dài nhất",
                f"Đáp án đa số là phương án dài nhất: {k1}/{n1} = {100*k1/n1:.0f}%; khoá là phương án dài nhất: "
                f"{k2}/{n2} = {100*k2/n2:.0f}%",
                "Tín hiệu yếu, chưa đủ để kết luận"])
    # 3. The same question (same words, same options) appears elsewhere in the raw release with another key?
    raw = [json.loads(line) for line in open(REPO_ROOT / "data/raw/data-processed-shuffled0.jsonl", encoding="utf-8")
           if line.strip()]
    index = collections.defaultdict(list)
    for r in raw:
        if len(r.get("options") or []) >= 2:
            index[_words(r["question"])].append(r)
    hits = 0
    for q in wrong:
        r = rows[q]
        opts = {_option_key(o) for o in r["options"]}
        key = _option_key(r["options"][r["answer_index"]])
        hits += any(o["id"] != q and {_option_key(x) for x in o["options"]} == opts
                    and _option_key(o["options"][o["answer_index"]]) != key for o in index[_words(r["question"])])
    out.append(["Câu y hệt (cùng chữ, cùng phương án) xuất hiện ở chỗ khác trong bản gốc với khoá khác",
                f"{hits}/{len(wrong)} câu",
                "Không có: các bản trùng có khoá mâu thuẫn đã được xử lý khi làm sạch. Các câu 'gần giống' tìm thấy "
                "đều khác số liệu hoặc khác thuốc, nên là câu khác nhau"])
    return out


def write_report(stats, n_wrong, n_test):
    out = ["# Phân tích các câu mà cả 7 model đều sai", "",
           "Sinh bởi `scripts/eval/build_allwrong_file.py`; cùng nội dung với sheet `TomTat_PhanTich` của "
           "`VM14K_125_cau_7_model_deu_sai.xlsx`.", "",
           f"**{n_wrong}/{n_test} câu test ({100 * n_wrong / n_test:.1f}%)** bị cả 7 model chọn sai. "
           "Mọi nhãn và nhóm gợi ý là tự động, chỉ để định hướng việc duyệt; kết luận cần ý kiến chuyên môn.", ""]
    for title, header, table, note in stats:
        out += [f"## {title}", "", "| " + " | ".join(map(str, header)) + " |",
                "|" + "---|" * len(header)]
        out += ["| " + " | ".join(str(c) for c in line) + " |" for line in table]
        if note:
            out += ["", f"*{note}*"]
        out.append("")
    REPORT.write_text("\n".join(out), encoding="utf-8")


if __name__ == "__main__":
    main()
