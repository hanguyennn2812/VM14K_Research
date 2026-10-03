"""Regression tests use copies or hand-built workbooks, never clinician files."""
import hashlib
import math
from pathlib import Path
import sys

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "eval"))
import analyze_review as a


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_inputs(directory, n=40, filled=False):
    directory.mkdir(parents=True, exist_ok=True)
    manifest = []
    for i in range(1, n + 1):
        manifest.append(dict(code=f"Q{i:03}", stratum="ngau_nhien" if i <= n // 2 else "uu_tien_test",
                             pool_size=1658 if i <= n // 2 else 100, priority_set="uu_tien_test" if i % 2 else "",
                             duyet_doi=int(i % 2 == 1)))
    pd.DataFrame(manifest).to_csv(directory / "sample_manifest.csv", index=False, encoding="utf-8")
    for filename, sheet, round_, reviewer in a.FILES:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        fields = a.ROUND1_FIELDS if round_ == 1 else a.ROUND2_FIELDS
        headers = [a.CODE, a.OPTIONS, "Đáp án chuẩn", "Đáp án đa số 7 model"] + [f[0] for f in fields]
        ws.append(headers)
        for item in manifest:
            if reviewer == 2 and not item["duyet_doi"]:
                continue
            row = {a.CODE: item["code"], a.OPTIONS: "A. Một\nB. Hai\nC. Ba\nD. Bốn",
                   "Đáp án chuẩn": "A", "Đáp án đa số 7 model": "B / C (3/7, hòa)"}
            if filled:
                if round_ == 1:
                    row.update({a.R1_ANSWER: "A", a.R1_STATUS: a.STATUS_CHOICES[0]})
                else:
                    row.update({a.R2_ANSWER: "A", a.KEY_STATUS: a.KEY_OK, a.R2_STATUS: a.STATUS_CHOICES[0]})
            ws.append([row.get(h) for h in headers])
        wb.save(directory / filename)
        wb.close()
    return directory


def edit(directory, filename, code, changes):
    wb = load_workbook(directory / filename)
    ws = wb["Vong2" if filename.startswith("vong2") else "Vong1"]
    columns = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        if row[0].value == code:
            for field, value in changes.items():
                ws.cell(row[0].row, columns[field]).value = value
    wb.save(directory / filename)
    wb.close()


def test_set_normalization():
    assert a.normalize_answer("A", n_options=4) == ("A",)
    assert a.normalize_answer(a.MULTI, "c,a / C; b", 4) == ("A", "B", "C")
    assert a.normalize_answer(a.NONE) == a.NONE
    assert a.normalize_answer(a.UNKNOWN) == a.UNKNOWN
    assert a.normalize_answer(None) is None
    assert a.normalize_answer(a.NONE) != a.normalize_answer(a.UNKNOWN)


@pytest.mark.parametrize("choice,letters,n", [
    (a.MULTI, "", 4), ("A", "AC", 4), (a.MULTI, "AE", 4),
    ("G", "", 4), ("mystery", "", 4), (a.MULTI, "A1", 4), (a.MULTI, "aa", 4),
])
def test_normalization_errors(choice, letters, n):
    with pytest.raises(ValueError):
        a.normalize_answer(choice, letters, n)


def test_validation_unknown_dropdown():
    row = {a.R2_ANSWER: "A", a.KEY_STATUS: "chắc đúng", a.R2_STATUS: "tốt"}
    answer, errors = a.validate_row(row, a.ROUND2_FIELDS, 4)
    assert answer == ("A",)
    assert len(errors) == 2


def test_wilson_zero_and_full():
    lo, hi = a.wilson(0, 60)
    assert lo == pytest.approx(0, abs=1e-15)
    assert hi == pytest.approx(.06017, abs=1e-5)
    assert a.wilson(60, 60) == pytest.approx((1-hi, 1))
    assert math.isnan(a.wilson(0, 0)[0])


def test_kappa_perfect_no_agreement_and_degenerate():
    assert a.cohen_kappa(["A", "B"] * 5, ["A", "B"] * 5) == 1
    assert a.cohen_kappa(["A", "B"] * 5, ["B", "A"] * 5) == -1
    assert math.isnan(a.cohen_kappa(["A"] * 5, ["A"] * 5))
    result = a.agreement([("A", "A"), ("B", "B"), (None, "A")], bootstraps=100)
    assert result["n"] == 2 and result["missing"] == 1
    assert result["kappa_ci"] == (1, 1)
    assert result["undefined_bootstraps"] > 0


def test_empty_report_does_not_touch_inputs(tmp_path):
    source = make_inputs(tmp_path / "empty", n=4)
    inputs = {p: digest(p) for p in source.iterdir()}
    result = a.analyze(source, bootstraps=20)
    assert "chưa có dữ liệu duyệt" in result["report"].read_text(encoding="utf-8")
    assert not result["problems"]
    assert result["summaries"]["ngau_nhien"]["missing"] == [2, 2, 2]
    assert all(digest(p) == h for p, h in inputs.items())


def test_simulation_full_analysis_and_determinism(tmp_path):
    source = make_inputs(tmp_path / "source")
    original = {p.name: digest(p) for p in source.iterdir()}
    copy = a.simulate(source, tmp_path / "sim1", seed=42)
    copy2 = a.simulate(source, tmp_path / "sim2", seed=42)
    first = a.analyze(copy, bootstraps=100)
    second = a.analyze(copy2, bootstraps=100)
    assert first["report"].read_text(encoding="utf-8") == second["report"].read_text(encoding="utf-8")
    assert first["problems"] and first["disagreements"]
    assert any(r.get("answer") == ("A", "B") for r in first["records"].values())
    assert any(r.get("answer") == a.UNKNOWN for r in first["records"].values())
    assert any(not r.get("filled") for r in first["records"].values())
    assert first["changes"][1]["changed"] > 0
    assert all(digest(source / name) == h for name, h in original.items())


def test_missing_duplicate_and_out_of_range_rows(tmp_path):
    source = make_inputs(tmp_path / "bad", n=6, filled=True)
    filename = "vong1_doc_lap.xlsx"
    edit(source, filename, "Q001", {a.R1_ANSWER: "E"})
    wb = load_workbook(source / filename)
    ws = wb["Vong1"]
    ws.append([c.value for c in ws[3]])  # duplicate Q002
    ws.delete_rows(4)  # remove Q003
    wb.save(source / filename)
    wb.close()
    result = a.analyze(source, bootstraps=20)
    messages = [(c, text) for c, _, text in result["problems"]]
    assert any(c == "Q001" and "vượt" in text for c, text in messages)
    assert any(c == "Q002" and "trùng" in text for c, text in messages)
    assert any(c == "Q003" and "thiếu dòng" in text for c, text in messages)
    assert result["summaries"]["ngau_nhien"]["n"] == 3


def test_adjudication_preserved_and_final_endpoints(tmp_path):
    source = make_inputs(tmp_path / "adjudicate", n=4, filled=True)
    edit(source, "vong2_doi_chieu_nguoi2.xlsx", "Q001", {a.KEY_STATUS: a.KEY_WRONG, a.R2_ANSWER: "B"})
    edit(source, "vong2_doi_chieu.xlsx", "Q002", {a.KEY_STATUS: a.KEY_UNKNOWN})
    first = a.analyze(source, bootstraps=20)
    assert first["final"]["Q001"]["flag"] == "chưa phân xử"
    assert first["summaries"]["ngau_nhien"]["u"] == 1
    wb = load_workbook(source / "phan_xu_template.xlsx")
    ws = wb["PhanXu"]
    cols = {c.value: c.column for c in ws[1]}
    for name, value in {a.KEY_STATUS: a.KEY_WRONG, a.R2_ANSWER: a.MULTI, a.R2_SET: "BA",
                        a.R2_STATUS: a.STATUS_CHOICES[-1], a.SOURCE: "Sách thử", a.ADJUDICATOR: "Người 3"}.items():
        ws.cell(2, cols[name], value)
    wb.save(source / "phan_xu.xlsx")
    wb.close()
    original = digest(source / "phan_xu.xlsx")
    final = a.analyze(source, bootstraps=20)
    assert digest(source / "phan_xu.xlsx") == original
    assert final["final"]["Q001"]["flag"] == "đã phân xử"
    assert final["final"]["Q001"]["answer"] == ("A", "B")
    summary = final["summaries"]["ngau_nhien"]
    assert summary["counts"] == [1, 1, 1]
    assert summary["sensitivity"] == (.5, 1)
    assert final["final"]["Q002"]["flag"] == "1 người duyệt"


def test_endpoint_union_and_majority_tie():
    assert a.endpoints({"key": a.KEY_MULTI, "status": a.STATUS_CHOICES[0]}) == (False, True, True)
    assert a.endpoints({"key": a.KEY_WRONG, "status": None}) == (True, None, True)
    assert a.endpoints({"key": None, "status": None}) == (None, None, None)
    ref = a.majority_reference({"Đáp án đa số 7 model": "B / C (3/7, hòa)"})
    assert ref == ("B", "C")
    assert a.change_direction(("A",), ("C",), ref) == "về phía mốc"


def test_manifest_duplicate_rejected_with_code(tmp_path):
    source = make_inputs(tmp_path / "duplicate", n=4)
    manifest = pd.read_csv(source / "sample_manifest.csv")
    manifest.loc[1, "code"] = "Q001"
    manifest.to_csv(source / "sample_manifest.csv", index=False)
    with pytest.raises(ValueError, match="Q001"):
        a.analyze(source)


def test_pool_formula_uses_random_members_and_pool_size(tmp_path):
    source = make_inputs(tmp_path / "pool", n=4, filled=True)
    # Q001 is a random member of S; Q003/Q004 are priority draws from pool 100.
    for code in ("Q001", "Q003", "Q004"):
        for filename in ("vong2_doi_chieu.xlsx", "vong2_doi_chieu_nguoi2.xlsx"):
            edit(source, filename, code, {a.KEY_STATUS: a.KEY_WRONG, a.R2_ANSWER: "B"})
    result = a.analyze(source, bootstraps=20)
    report = result["report"].read_text(encoding="utf-8")
    assert "| uu_tien_test | E3 | 1 | 100 × 2/2 | 101.00 | 0 / 0 |" in report


def test_incomplete_double_review_is_missing_not_adjudicated(tmp_path):
    source = make_inputs(tmp_path / "incomplete", n=4, filled=True)
    edit(source, "vong2_doi_chieu_nguoi2.xlsx", "Q001", {a.R2_ANSWER: None})
    result = a.analyze(source, bootstraps=20)
    assert not result["disagreements"]
    assert result["final"]["Q001"]["flag"] == "thiếu duyệt đôi"
    assert result["final"]["Q001"]["key"] is None




def test_invalid_adjudication_cannot_resolve_disagreement(tmp_path):
    source = make_inputs(tmp_path / "invalid_decision", n=4, filled=True)
    edit(source, "vong2_doi_chieu_nguoi2.xlsx", "Q001", {a.R2_ANSWER: "B", a.KEY_STATUS: a.KEY_WRONG})
    a.analyze(source, bootstraps=20)
    wb = load_workbook(source / "phan_xu_template.xlsx")
    ws = wb["PhanXu"]
    cols = {c.value: c.column for c in ws[1]}
    ws.cell(2, cols[a.R2_ANSWER], a.MULTI)  # no letters, no decision metadata
    wb.save(source / "phan_xu.xlsx")
    wb.close()
    result = a.analyze(source, bootstraps=20)
    assert result["final"]["Q001"]["flag"] == "chưa phân xử"
    assert any(file == "phan_xu.xlsx" and code == "Q001" for code, file, _ in result["problems"])


def test_blank_flag_with_status_counts_as_absent(tmp_path):
    source = make_inputs(tmp_path / "flags", n=4, filled=True)  # status "Bình thường", flags blank
    result = a.analyze(source, bootstraps=20)
    metric = result["agreements"]["tất cả", a.DEFECT_FLAGS[0][0]]
    assert metric["n"] == 2 and metric["missing"] == 0 and metric["raw"] == 1
    assert math.isnan(metric["kappa"])  # one category only: kappa undefined


def test_blank_flag_without_status_is_missing(tmp_path):
    source = make_inputs(tmp_path / "flags_missing", n=4, filled=False)
    edit(source, "vong1_doc_lap.xlsx", "Q001", {a.R1_ANSWER: "A"})
    edit(source, "vong1_doc_lap_nguoi2.xlsx", "Q001", {a.R1_ANSWER: "A"})
    result = a.analyze(source, bootstraps=20)
    metric = result["agreements"]["tất cả", a.DEFECT_FLAGS[0][0]]
    assert metric["n"] == 0


def test_flag_status_contradiction_is_warned_not_dropped(tmp_path):
    source = make_inputs(tmp_path / "flags_conflict", n=4, filled=True)
    edit(source, "vong1_doc_lap.xlsx", "Q001", {a.DEFECT_FLAGS[0][0]: "Có"})  # status stays "Bình thường"
    result = a.analyze(source, bootstraps=20)
    assert any(code == "Q001" and msg.startswith("cảnh báo") for code, _, msg in result["problems"])
    assert result["records"][1, 1, "Q001"]["valid"]


def test_status_only_disagreement_goes_to_adjudication(tmp_path):
    source = make_inputs(tmp_path / "status_only", n=4, filled=True)
    edit(source, "vong2_doi_chieu_nguoi2.xlsx", "Q001", {a.R2_STATUS: a.STATUS_CHOICES[-1],
                                                         a.DEFECT_FLAGS[2][0]: "Có"})
    result = a.analyze(source, bootstraps=20)
    assert result["final"]["Q001"]["flag"] == "chưa phân xử"
    assert "Q001" in {code for code, _, _ in result["disagreements"]}
